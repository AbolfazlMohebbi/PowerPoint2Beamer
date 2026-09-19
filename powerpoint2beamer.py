#!/usr/bin/env python3
"""PowerPoint2Beamer: convert a PowerPoint (.pptx) presentation into a simple Beamer LaTeX project.

Output:
    OUTDIR/main.tex                     preamble + \\input of every slide
    OUTDIR/slides/slide_<i>.tex         one frame per slide
    OUTDIR/images/slide_<i>_image_<j>.* pictures and video poster frames

Usage:
    python powerpoint2beamer.py deck.pptx [-o OUTDIR] [--theme Madrid] [--color auto]

Author:
    Abolfazl Mohebbi, PhD., P. Eng.
    Professor in Mechanical and Biomedical Engineering
    Polytechnique Montréal
    abolfazl.mohebbi@polymtl.ca

License: MIT (see LICENSE)
"""

import argparse
import colorsys
import copy
import io
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.shapetree import SlideShapeFactory

try:
    from PIL import Image, ImageChops
except ImportError:  # Pillow is optional; only needed for format conversion/cropping
    Image = None


# --------------------------------------------------------------------------
# Themes
# --------------------------------------------------------------------------

THEMES = [
    "default", "AnnArbor", "Antibes", "Bergen", "Berkeley", "Berlin", "Boadilla",
    "CambridgeUS", "Copenhagen", "Darmstadt", "Dresden", "Frankfurt", "Goettingen",
    "Hannover", "Ilmenau", "JuanLesPins", "Luebeck", "Madrid", "Malmoe", "Marburg",
    "Montpellier", "PaloAlto", "Pittsburgh", "Rochester", "Singapore", "Szeged",
    "Warsaw", "metropolis",
]

# Beamer colour themes with the (approximate) dominant colour they paint the
# headline/frame title with, and their page background colour.
COLOR_THEMES = {
    "whale":     ((51, 51, 179),   (255, 255, 255)),
    "dolphin":   ((140, 140, 215), (255, 255, 255)),
    "seahorse":  ((200, 200, 235), (255, 255, 255)),
    "beaver":    ((170, 0, 0),     (255, 255, 255)),
    "crane":     ((255, 191, 0),   (255, 255, 255)),
    "wolverine": ((255, 204, 51),  (255, 255, 255)),
    "spruce":    ((0, 90, 50),     (255, 255, 255)),
    "seagull":   ((150, 150, 150), (255, 255, 255)),
    "dove":      ((60, 60, 60),    (255, 255, 255)),
    "beetle":    ((51, 51, 153),   (153, 153, 153)),
    "fly":       ((220, 220, 220), (77, 77, 77)),
    "albatross": ((150, 150, 255), (20, 20, 60)),
}
OTHER_COLOR_THEMES = ["default", "orchid", "rose", "lily", "structure", "sidebartab"]


# --------------------------------------------------------------------------
# LaTeX escaping
# --------------------------------------------------------------------------

LATEX_ESCAPES = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
    "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}", "<": r"\textless{}", ">": r"\textgreater{}",
    # common Unicode symbols pdflatex does not know out of the box
    "\u2022": r"\textbullet{}", "\u2013": "--", "\u2014": "---", "\u2026": r"\ldots{}",
    "\u2192": r"$\rightarrow$", "\u2190": r"$\leftarrow$", "\u21d2": r"$\Rightarrow$",
    "\u2264": r"$\leq$", "\u2265": r"$\geq$", "\u2260": r"$\neq$", "\u00d7": r"$\times$",
    "\u00b1": r"$\pm$", "\u2248": r"$\approx$", "\u221e": r"$\infty$", "\u2211": r"$\sum$",
    "\u03b1": r"$\alpha$", "\u03b2": r"$\beta$", "\u03b3": r"$\gamma$", "\u03b4": r"$\delta$",
    "\u03b5": r"$\epsilon$", "\u03b8": r"$\theta$", "\u03bb": r"$\lambda$", "\u03bc": r"$\mu$",
    "\u03c0": r"$\pi$", "\u03c3": r"$\sigma$", "\u03c6": r"$\phi$", "\u03c9": r"$\omega$",
    "\u0394": r"$\Delta$", "\u03a3": r"$\Sigma$", "\u03a9": r"$\Omega$",
    "\u2713": r"$\checkmark$", "\u00a0": "~", "\u200b": "",
}


def esc(text):
    out = []
    for ch in text:
        if 0xF000 <= ord(ch) <= 0xF0FF:  # Symbol/Wingdings private-use glyphs
            continue
        if 0x1D400 <= ord(ch) <= 0x1D7FF or ch in "ℎℒ":  # math letters (𝑠, ℒ, ...)
            out.append("$" + _math_char(ch).strip() + "$")
            continue
        if ch in GREEK and ch not in LATEX_ESCAPES:  # Greek letters in normal text: ζ -> $\zeta$
            out.append("$\\" + GREEK[ch] + "$")
            continue
        out.append(LATEX_ESCAPES.get(ch, ch))
    return "".join(out)


def esc_url(url):
    return url.replace("\\", "/").replace("%", r"\%").replace("#", r"\#")


# --------------------------------------------------------------------------
# Colours
# --------------------------------------------------------------------------

def _lab(rgb):
    def lin(c):
        c /= 255.0
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (lin(c) for c in rgb)
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883

    def f(t):
        return t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    fx, fy, fz = f(x), f(y), f(z)
    return 116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz)


def color_distance(a, b):
    return math.dist(_lab(a), _lab(b))


def hex_rgb(val):
    return tuple(int(val[i:i + 2], 16) for i in (0, 2, 4))


def deck_colors(prs):
    """Return (accent, background) RGB tuples of the deck's theme."""
    master = prs.slide_master
    scheme = {}
    try:
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        for el in theme.find(".//" + qn("a:clrScheme")):
            name = etree.QName(el).localname
            srgb = el.find(qn("a:srgbClr"))
            sysc = el.find(qn("a:sysClr"))
            if srgb is not None:
                scheme[name] = hex_rgb(srgb.get("val"))
            elif sysc is not None and sysc.get("lastClr"):
                scheme[name] = hex_rgb(sysc.get("lastClr"))
    except (KeyError, TypeError, AttributeError):
        pass

    clr_map = master._element.find(qn("p:clrMap"))
    aliases = dict(clr_map.attrib) if clr_map is not None else {}
    aliases.setdefault("bg1", "lt1")
    aliases.setdefault("tx1", "dk1")
    aliases.setdefault("bg2", "lt2")
    aliases.setdefault("tx2", "dk2")
    THEME.clear()
    THEME.update(scheme)
    THEME.update({alias: scheme[name] for alias, name in aliases.items() if name in scheme})

    def resolve(fill_parent):
        if fill_parent is None:
            return None
        srgb = fill_parent.find(".//" + qn("a:srgbClr"))
        if srgb is not None:
            return hex_rgb(srgb.get("val"))
        sch = fill_parent.find(".//" + qn("a:schemeClr"))
        if sch is not None:
            name = sch.get("val")
            return scheme.get(aliases.get(name, name))
        return None

    # background: first slide that is not a title slide, then its layout, then master
    background = None
    sample = next((s for s in prs.slides if not _is_title_slide(s)), None)
    candidates = [sample, sample.slide_layout] if sample is not None else []
    for obj in candidates + [master]:
        bg = obj._element.find(qn("p:cSld")).find(qn("p:bg"))
        if bg is not None:
            background = resolve(bg)
            if background:
                break
    background = background or scheme.get("lt1", (255, 255, 255))
    accent = scheme.get("accent1") or scheme.get("dk2") or (51, 51, 179)
    return accent, background


THEME = {}  # colour scheme of the deck being converted: name (accent1, tx1, ...) -> RGB


def run_color(rpr):
    """The colour a text run was given, or None."""
    fill = rpr.find(qn("a:solidFill")) if rpr is not None else None
    if fill is None:
        return None
    srgb, scheme = fill.find(qn("a:srgbClr")), fill.find(qn("a:schemeClr"))
    if srgb is not None:
        return hex_rgb(srgb.get("val"))
    if scheme is not None:
        return THEME.get(scheme.get("val"))
    return None


def color_role(rgb):
    """None (normal text colour), 'alert' (red emphasis) or 'rgb' (keep the colour)."""
    if rgb is None or rgb == THEME.get("tx1"):
        return None
    hue, light, sat = colorsys.rgb_to_hls(*(c / 255 for c in rgb))
    if light < 0.15 or light > 0.92 or sat < 0.25:
        return None  # black, grey or white: the normal text colour
    if (hue * 360 <= 25 or hue * 360 >= 335) and sat >= 0.5:
        return "alert"
    return "rgb"


def pick_color_theme(accent, background):
    def score(item):
        struct, bg = item[1]
        return color_distance(accent, struct) + color_distance(background, bg)
    return min(COLOR_THEMES.items(), key=score)[0]


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------

CONTENT_TYPE_EXT = {
    "image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif",
    "image/bmp": "bmp", "image/tiff": "tiff", "image/x-emf": "emf", "image/emf": "emf",
    "image/x-wmf": "wmf", "image/wmf": "wmf", "image/svg+xml": "svg",
    "application/pdf": "pdf",
}


def _convert_vector(blob, ext, out_png):
    """Convert EMF/WMF to PNG. Returns True on success."""
    if Image is not None:
        try:
            im = Image.open(io.BytesIO(blob))
            im.load(dpi=200)
            im.save(out_png)
            return True
        except Exception:
            pass
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in." + ext)
        with open(src, "wb") as f:
            f.write(blob)
        cmds = []
        if shutil.which("inkscape"):
            cmds.append(["inkscape", src, "--export-type=png",
                         "--export-filename=" + os.path.join(tmp, "in.png")])
        soffice = shutil.which("soffice") or shutil.which("soffice.exe")
        if soffice:
            cmds.append([soffice, "--headless", "--convert-to", "png", "--outdir", tmp, src])
        for cmd in cmds:
            try:
                subprocess.run(cmd, capture_output=True, timeout=120)
            except (OSError, subprocess.SubprocessError):
                continue
            produced = os.path.join(tmp, "in.png")
            if os.path.exists(produced):
                shutil.move(produced, out_png)
                return True
    return False


def save_image(blob, content_type, images_dir, name, crop=(0, 0, 0, 0)):
    """Save an image blob as images_dir/name.<ext> in a pdflatex-friendly format.

    Returns the file name that was written, and whether pdflatex can include it.
    """
    ext = CONTENT_TYPE_EXT.get(content_type, "png")
    cropped = any(c > 0 for c in crop)

    if ext in ("png", "jpg", "pdf") and not cropped:
        fname = f"{name}.{ext}"
        with open(os.path.join(images_dir, fname), "wb") as f:
            f.write(blob)
        return fname, True

    if ext in ("emf", "wmf"):
        fname = f"{name}.png"
        if _convert_vector(blob, ext, os.path.join(images_dir, fname)):
            return fname, True
        fname = f"{name}.{ext}"
        with open(os.path.join(images_dir, fname), "wb") as f:
            f.write(blob)
        return fname, False

    if Image is not None and ext != "svg":
        try:
            im = Image.open(io.BytesIO(blob))
            if cropped:
                w, h = im.size
                l, t, r, b = crop
                im = im.crop((int(l * w), int(t * h), int(w - r * w), int(h - b * h)))
            out_ext = "jpg" if ext == "jpg" else "png"
            if out_ext == "png" and im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
                im = im.convert("RGBA")
            fname = f"{name}.{out_ext}"
            im.save(os.path.join(images_dir, fname))
            return fname, True
        except Exception:
            pass

    fname = f"{name}.{ext}"
    with open(os.path.join(images_dir, fname), "wb") as f:
        f.write(blob)
    return fname, ext in ("png", "jpg", "pdf")


# --------------------------------------------------------------------------
# Office equations (OMML) -> LaTeX math
# --------------------------------------------------------------------------

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

MATH_SYMBOLS = {
    "−": "-", "∗": "*", "×": r"\times ", "·": r"\cdot ", "⋅": r"\cdot ",
    "÷": r"\div ", "±": r"\pm ", "∓": r"\mp ", "∞": r"\infty ",
    "→": r"\rightarrow ", "←": r"\leftarrow ", "⇒": r"\Rightarrow ",
    "⇔": r"\Leftrightarrow ", "↔": r"\leftrightarrow ", "↦": r"\mapsto ",
    "≤": r"\leq ", "≥": r"\geq ", "≠": r"\neq ", "≈": r"\approx ",
    "≡": r"\equiv ", "∝": r"\propto ", "∼": r"\sim ", "≪": r"\ll ", "≫": r"\gg ",
    "∈": r"\in ", "∉": r"\notin ", "⊂": r"\subset ", "⊆": r"\subseteq ",
    "∪": r"\cup ", "∩": r"\cap ", "∀": r"\forall ", "∃": r"\exists ",
    "∂": r"\partial ", "∇": r"\nabla ", "…": r"\ldots ", "⋯": r"\cdots ",
    "⋮": r"\vdots ", "⋱": r"\ddots ", "′": "'", "″": "''", "°": r"^\circ ",
    "∘": r"\circ ", "∑": r"\sum ", "∏": r"\prod ", "∫": r"\int ",
    "√": r"\sqrt ", "‖": r"\| ", "⟨": r"\langle ", "⟩": r"\rangle ",
    "ℏ": r"\hbar ", "ℓ": r"\ell ", "∅": r"\emptyset ", "∧": r"\wedge ",
    "∨": r"\vee ", "¬": r"\neg ", "∥": r"\parallel ", "⊥": r"\perp ",
    " ": " ", " ": r"\, ", "​": "", "⁡": "", "⁢": "", "⁣": "",
    "{": r"\{", "}": r"\}", "%": r"\%", "#": r"\#", "$": r"\$", "_": r"\_", "~": r"\sim ",
    "\\": r"\backslash ",
}
GREEK = {
    "α": "alpha", "β": "beta", "γ": "gamma", "δ": "delta", "ε": "epsilon",
    "ϵ": "epsilon", "ζ": "zeta", "η": "eta", "θ": "theta", "ϑ": "vartheta",
    "ι": "iota", "κ": "kappa", "λ": "lambda", "μ": "mu", "ν": "nu",
    "ξ": "xi", "π": "pi", "ρ": "rho", "σ": "sigma", "ς": "varsigma",
    "τ": "tau", "υ": "upsilon", "φ": "phi", "ϕ": "phi", "χ": "chi",
    "ψ": "psi", "ω": "omega", "Γ": "Gamma", "Δ": "Delta", "Θ": "Theta",
    "Λ": "Lambda", "Ξ": "Xi", "Π": "Pi", "Σ": "Sigma", "Υ": "Upsilon",
    "Φ": "Phi", "Ψ": "Psi", "Ω": "Omega",
}
FUNCTIONS = {"sin", "cos", "tan", "cot", "sec", "csc", "arcsin", "arccos", "arctan", "sinh",
             "cosh", "tanh", "exp", "log", "ln", "lg", "lim", "max", "min", "sup", "inf", "det",
             "arg", "deg", "dim", "gcd", "ker", "Pr"}
ACCENTS = {"̂": r"\hat", "̃": r"\tilde", "̄": r"\bar", "̅": r"\bar",
           "̇": r"\dot", "̈": r"\ddot", "⃛": r"\dddot", "⃗": r"\vec",
           "̌": r"\check", "̆": r"\breve", "́": r"\acute", "̀": r"\grave"}
NARY = {"∫": r"\int", "∬": r"\iint", "∭": r"\iiint", "∮": r"\oint",
        "∑": r"\sum", "∏": r"\prod", "∐": r"\coprod", "⋃": r"\bigcup",
        "⋂": r"\bigcap"}
DELIMS = {"": ".", "{": r"\{", "}": r"\}", "‖": r"\|", "⟨": r"\langle",
          "⟩": r"\rangle", "⌊": r"\lfloor", "⌋": r"\rfloor", "⌈": r"\lceil",
          "⌉": r"\rceil"}


def _math_char(ch):
    """One Unicode character -> LaTeX math (handles 'Mathematical Italic' letters etc.)."""
    if ch in MATH_SYMBOLS:
        return MATH_SYMBOLS[ch]
    if ch in GREEK:
        return "\\" + GREEK[ch] + " "
    if ord(ch) < 128:
        return ch
    base = unicodedata.normalize("NFKC", ch)
    if base == ch or len(base) != 1:
        return ch
    name = unicodedata.name(ch, "")
    inner = _math_char(base)
    if "SCRIPT" in name or "ℒ" == ch:
        return r"\mathcal{%s}" % inner.strip()
    if "DOUBLE-STRUCK" in name:
        return r"\mathbb{%s}" % inner.strip()
    if "BOLD" in name:
        return r"\boldsymbol{%s}" % inner.strip() if inner.startswith("\\") else r"\mathbf{%s}" % inner
    return inner


def _mval(el, tag, default=None):
    """Value of a property like <m:chr m:val="x"/> inside el (e.g. inside m:dPr)."""
    if el is None:
        return default
    found = el.find(".//" + M + tag)
    if found is None:
        return default
    return found.get(M + "val", default if default is not None else "")


def _m(el):
    """Convert the children of an OMML element to LaTeX."""
    if el is None:
        return ""
    return "".join(_m_node(child) for child in el)


def _arg(el):
    s = _m(el).strip()
    return s if len(s) == 1 and s.isalnum() else "{" + s + "}"


def _m_node(el):
    tag = etree.QName(el).localname
    if tag.endswith("Pr") or tag == "ctrlPr":
        return ""
    if tag == "r":
        text = "".join(t.text or "" for t in el.iter(M + "t"))
        rpr = el.find(M + "rPr")
        plain = rpr is not None and (rpr.find(M + "nor") is not None or _mval(rpr, "sty") == "p")
        if text in FUNCTIONS:
            return "\\" + text + " "
        if plain and len(text) > 1 and text.isascii() and text.isalpha():
            return r"\mathrm{%s}" % text
        if any(c.isalpha() and 0x80 <= ord(c) <= 0x24F for c in text):  # é, à, ... in math
            return r"\text{%s}" % esc(text)
        return "".join(_math_char(c) for c in text)
    if tag == "f":
        num, den = _m(el.find(M + "num")), _m(el.find(M + "den"))
        if _mval(el.find(M + "fPr"), "type") == "lin":
            return "%s/%s" % (_arg(el.find(M + "num")), _arg(el.find(M + "den")))
        return r"\frac{%s}{%s}" % (num.strip(), den.strip())
    if tag == "sSup":
        sup = _m(el.find(M + "sup")).strip()
        if sup and set(sup) == {"'"}:  # f′ is written f' in LaTeX, not f^'
            return _arg(el.find(M + "e")) + sup
        return "%s^%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sup")))
    if tag == "sSub":
        return "%s_%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sub")))
    if tag == "sSubSup":
        return "%s_%s^%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sub")), _arg(el.find(M + "sup")))
    if tag == "sPre":
        return "{}_%s^%s%s" % (_arg(el.find(M + "sub")), _arg(el.find(M + "sup")), _arg(el.find(M + "e")))
    if tag == "d":
        pr = el.find(M + "dPr")
        beg, end = _mval(pr, "begChr", "("), _mval(pr, "endChr", ")")
        sep = _mval(pr, "sepChr", "|")
        parts = [_m(e).strip() for e in el.findall(M + "e")]
        inner = (" %s " % MATH_SYMBOLS.get(sep, sep).strip()).join(parts)
        return r"\left%s %s \right%s " % (DELIMS.get(beg, beg), inner, DELIMS.get(end, end))
    if tag == "rad":
        deg = _m(el.find(M + "deg")).strip()
        body = _m(el.find(M + "e")).strip()
        return r"\sqrt[%s]{%s}" % (deg, body) if deg else r"\sqrt{%s}" % body
    if tag == "acc":
        cmd = ACCENTS.get(_mval(el.find(M + "accPr"), "chr", "̂"), r"\hat")
        return "%s{%s}" % (cmd, _m(el.find(M + "e")).strip())
    if tag == "bar":
        cmd = r"\underline" if _mval(el.find(M + "barPr"), "pos") == "bot" else r"\overline"
        return "%s{%s}" % (cmd, _m(el.find(M + "e")).strip())
    if tag == "groupChr":
        pr = el.find(M + "groupChrPr")
        chr_ = _mval(pr, "chr", "⏟")
        body = _m(el.find(M + "e")).strip()
        if chr_ == "⏟":
            return r"\underbrace{%s}" % body
        if chr_ == "⏞":
            return r"\overbrace{%s}" % body
        cmd = r"\overset" if _mval(pr, "pos") == "top" else r"\underset"
        return r"%s{%s}{%s}" % (cmd, _math_char(chr_).strip(), body)
    if tag in ("limLow", "limUpp"):
        base = _m(el.find(M + "e")).strip()
        lim = _m(el.find(M + "lim")).strip()
        if base.startswith("\\") and base[1:].strip() in FUNCTIONS:
            return "%s%s{%s} " % (base.strip(), "_" if tag == "limLow" else "^", lim)
        cmd = r"\underset" if tag == "limLow" else r"\overset"
        return r"%s{%s}{%s}" % (cmd, lim, base)
    if tag == "func":
        name = _m(el.find(M + "fName")).strip()
        if name.startswith(r"\mathrm{") and name[8:-1] in FUNCTIONS:
            name = "\\" + name[8:-1]
        return "%s %s" % (name, _m(el.find(M + "e")).strip())
    if tag == "nary":
        pr = el.find(M + "naryPr")
        op = NARY.get(_mval(pr, "chr", "∫"), r"\int")
        s = op
        if _mval(pr, "subHide") not in ("1", "on", "true"):
            sub = _m(el.find(M + "sub")).strip()
            s += "_{%s}" % sub if sub else ""
        if _mval(pr, "supHide") not in ("1", "on", "true"):
            sup = _m(el.find(M + "sup")).strip()
            s += "^{%s}" % sup if sup else ""
        return s + " " + _m(el.find(M + "e")).strip()
    if tag == "eqArr":
        rows = [_m(e).strip() for e in el.findall(M + "e")]
        if not any("&" in r for r in rows):
            rows = [r.replace("=", "&=", 1) for r in rows]
        return "\\begin{aligned}" + r" \\ ".join(rows) + r"\end{aligned}"
    if tag == "m":
        rows = [" & ".join(_m(e).strip() for e in mr.findall(M + "e"))
                for mr in el.findall(M + "mr")]
        return "\\begin{matrix}" + r" \\ ".join(rows) + r"\end{matrix}"
    # m:e, m:box, m:borderBox, m:phant, m:num, ... : just their content
    return _m(el)


def _tidy_math(s):
    s = " ".join(s.split())
    s = s.replace("{ ", "{").replace(" }", "}")
    # Entr\text{é}e -> \text{Entrée}
    return re.sub(r"([A-Za-z]*)\\text\{([^{}\\]*)\}([A-Za-z]*)", r"\\text{\1\2\3}", s)


def omml_to_latex(a14m):
    """Convert an a14:m element to a list of LaTeX strings ($...$ or \\[...\\])."""
    out = []
    for el in a14m:
        tag = etree.QName(el).localname
        if tag == "oMathPara":
            for om in el.iter(M + "oMath"):
                out.append(r"\[ " + _tidy_math(_m(om)) + r" \]")
        elif tag == "oMath":
            out.append("$" + _tidy_math(_m(el)) + "$")
    return out


# --------------------------------------------------------------------------
# Slide analysis helpers
# --------------------------------------------------------------------------

SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
BULLET_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, None}


def _ph_type(shape):
    if not shape.is_placeholder:
        return "none"
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def _is_title_slide(slide):
    return any(_ph_type(s) == PP_PLACEHOLDER.CENTER_TITLE for s in slide.placeholders)


def _bullet_from_ppr(ppr):
    if ppr is None:
        return None
    if ppr.find(qn("a:buNone")) is not None:
        return "none"
    if ppr.find(qn("a:buAutoNum")) is not None:
        return "enumerate"
    if ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buBlip")) is not None:
        return "itemize"
    return None


def _inherited_bullet(shape, level):
    """Look up the bullet style in the shape's list style and its layout/master placeholders."""
    chain = [shape]
    base = getattr(shape, "_base_placeholder", None)
    while base is not None:
        chain.append(base)
        base = getattr(base, "_base_placeholder", None)
    for sh in chain:
        lst = sh._element.find(".//" + qn("a:lstStyle"))
        if lst is not None:
            kind = _bullet_from_ppr(lst.find(qn(f"a:lvl{level + 1}pPr")))
            if kind:
                return kind
    return None


def paragraph_kind(shape, paragraph):
    kind = _bullet_from_ppr(paragraph._p.pPr)
    if kind is None and shape.is_placeholder:
        kind = _inherited_bullet(shape, paragraph.level)
        if kind is None and _ph_type(shape) in BULLET_PLACEHOLDERS:
            kind = "itemize"
    return None if kind in (None, "none") else kind


def _paragraph_children(p):
    """Yield the run-level children of a paragraph, unwrapping mc:AlternateContent."""
    for el in p:
        if etree.QName(el).localname == "AlternateContent":
            choice = el.find(MC + "Choice")
            if choice is not None:
                yield from _paragraph_children(choice)
        else:
            yield el


def paragraph_latex(paragraph, part, line_break=r"\\ "):
    """Paragraph text with bold/italic/colour/hyperlinks/equations, equal runs merged.

    Equations are always black: colours are only applied to ordinary text runs.
    """
    pieces = []  # [text, bold, italic, url, color] for text, or a raw LaTeX string
    for el in _paragraph_children(paragraph._p):
        tag = etree.QName(el).localname
        if tag == "br":
            pieces.append(line_break)
            continue
        if tag == "m":  # a14:m wraps an Office equation
            pieces.extend(omml_to_latex(el))
            continue
        if tag not in ("r", "fld"):
            continue
        text = "".join(t.text or "" for t in el.iter(qn("a:t")))
        rpr = el.find(qn("a:rPr"))
        bold = rpr is not None and rpr.get("b") in ("1", "true")
        italic = rpr is not None and rpr.get("i") in ("1", "true")
        url = None
        link = rpr.find(qn("a:hlinkClick")) if rpr is not None else None
        if link is not None and link.get(qn("r:id")):
            try:
                url = part.rels[link.get(qn("r:id"))].target_ref
            except KeyError:
                url = None
        rgb = run_color(rpr)
        color = (color_role(rgb), rgb) if color_role(rgb) else None
        if color and color[0] == "alert":
            color = ("alert", None)  # all reds become the same \alert
        if pieces and isinstance(pieces[-1], list) and pieces[-1][1:] == [bold, italic, url, color]:
            pieces[-1][0] += text
        else:
            pieces.append([text, bold, italic, url, color])

    out = []
    for piece in pieces:
        if isinstance(piece, str):
            out.append(piece)
            continue
        text, bold, italic, url, color = piece
        s = esc(text)
        if s.strip():
            if bold:
                s = r"\textbf{" + s + "}"
            if italic:
                s = r"\textit{" + s + "}"
            if color and color[0] == "alert":
                s = r"\alert{" + s + "}"
            elif color:
                s = r"\textcolor[RGB]{%d,%d,%d}{%s}" % (*color[1], s)
            if url:
                s = r"\href{" + esc_url(url) + "}{" + s + "}"
        out.append(s)
    return "".join(out).strip()


def render_text_frame(shape, part, builds=None):
    """Turn a text frame into itemize/enumerate blocks and plain paragraphs.

    builds: {paragraph index: click} for paragraphs animated one by one.
    """
    lines, stack = [], []  # stack of (env, level)
    builds = builds or {}

    def close_to(level):
        while stack and stack[-1][1] > level:
            lines.append("  " * (len(stack) - 1) + r"\end{" + stack.pop()[0] + "}")

    for index, p in enumerate(shape.text_frame.paragraphs):
        text = paragraph_latex(p, part)
        if not text:
            continue
        click = builds.get(index)
        kind = paragraph_kind(shape, p)
        if kind is None:
            close_to(-1)
            if lines:
                lines.append("")
            lines.append(r"\uncover<%d->{%s}" % (click, text) if click else text)
            continue
        level = min(p.level, 3)
        close_to(level)
        if stack and stack[-1][1] == level and stack[-1][0] != kind:
            close_to(level - 1)
        if not stack or stack[-1][1] < level:
            lines.append("  " * len(stack) + r"\begin{" + kind + "}")
            stack.append((kind, level))
        item = r"\item<%d-> " % click if click else r"\item "
        lines.append("  " * len(stack) + item + text)
    close_to(-1)
    return "\n".join(lines)


def render_table(table, part):
    ncols = len(table.columns)
    rows = [r"\begin{tabular}{|" + "l|" * ncols + "}", r"\hline"]
    for row in table.rows:
        cells, skip = [], 0
        for cell in row.cells:
            if skip:
                skip -= 1
                continue
            text = " ".join(filter(None, (paragraph_latex(p, part, " ")
                                          for p in cell.text_frame.paragraphs)))
            text = re.sub(r"\\\[ ?(.*?) ?\\\]", r"$\1$", text)  # no display math in a cell
            if cell.is_spanned:
                text = ""
            if cell.is_merge_origin and cell.span_width > 1:
                skip = cell.span_width - 1
                text = r"\multicolumn{%d}{|l|}{%s}" % (cell.span_width, text)
            cells.append(text)
        rows.append(" & ".join(cells) + r" \\ \hline")
    rows.append(r"\end{tabular}")
    body = "\n".join(rows)
    # widest cell text per column, to see whether the table fits the slide width
    widths = [0] * ncols
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < ncols:
                widths[i] = max(widths[i], len(cell.text_frame.text))
    if ncols > 5 or sum(widths) + 3 * ncols > 60:
        body = "\\resizebox{\\linewidth}{!}{%\n" + body + "}"
    return body


# --------------------------------------------------------------------------
# Animations, sections, language, charts
# --------------------------------------------------------------------------

def read_animations(slide):
    """Click animations -> ({shape id: (enter click, exit click)}, {(shape id, paragraph): click}).

    Beamer overlay 1 is the slide before any click, so the first click is overlay 2.
    Only appear/disappear effects matter; emphasis and motion paths are ignored.
    """
    shapes, paragraphs = {}, {}
    seq = next((c for c in slide._element.iter(qn("p:cTn")) if c.get("nodeType") == "mainSeq"), None)
    if seq is None or seq.find(qn("p:childTnLst")) is None:
        return shapes, paragraphs
    click = 1
    for step in seq.find(qn("p:childTnLst")):
        effects = [c for c in step.iter(qn("p:cTn")) if c.get("presetClass")]
        if effects and effects[0].get("nodeType") == "clickEffect":
            click += 1
        for ctn in effects:
            cls = ctn.get("presetClass")
            target = ctn.find(".//" + qn("p:spTgt"))
            if cls not in ("entr", "exit") or target is None:
                continue
            spid = int(target.get("spid"))
            rng = target.find(".//" + qn("p:pRg"))
            if rng is not None and cls == "entr":
                for i in range(int(rng.get("st")), int(rng.get("end")) + 1):
                    paragraphs.setdefault((spid, i), click)
                continue
            enter, leave = shapes.get(spid, (None, None))
            if cls == "entr" and enter is None:
                enter = click
            elif cls == "exit":
                leave = click
            shapes[spid] = (enter, leave)
    return shapes, paragraphs


def overlay_spec(enter, leave):
    """<3->, <-4>, <3-4> or None (always visible)."""
    first = enter if enter and enter > 1 else None
    last = leave - 1 if leave else None
    if first is None and last is None:
        return None
    if last is not None and last < (first or 1):
        return "<%d>" % (first or 1)
    return "<%s-%s>" % (first or "", last or "")


def slide_title(slide):
    shape = slide.shapes.title
    if shape is None or not shape.has_text_frame:
        return ""
    return " ".join(shape.text_frame.text.replace("\v", " ").split())


def find_sections(prs, mode="auto"):
    """{slide index: section name}. Uses PowerPoint sections, else 'Section Header' slides,
    else numbered titles ('2.1 Équations du mouvement' starts 'Équations du mouvement')."""
    if mode == "none":
        return {}
    p14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
    ids = [int(s.get("id")) for s in prs.slides._sldIdLst]
    found = {}
    for sec in prs.part._element.iter(p14 + "section"):
        members = [int(s.get("id")) for s in sec.iter(p14 + "sldId")]
        if members and members[0] in ids:
            found[ids.index(members[0])] = sec.get("name", "")
    if len(found) >= 2 or mode == "pptx":
        return found if len(found) >= 2 else {}

    found = {i: slide_title(s) for i, s in enumerate(prs.slides)
             if s.slide_layout._element.get("type") == "secHead" and slide_title(s)}
    if len(found) >= 2:
        return found

    found, current = {}, None
    for i, slide in enumerate(prs.slides):
        m = re.match(r"^\s*(\d+(?:\.\d+)+)\.?\s+(.+)", slide_title(slide))
        if m and m.group(1) != current:
            current = m.group(1)
            found[i] = m.group(2)
    return found if len(found) >= 2 else {}


BABEL = {"fr": "french", "en": "english", "de": "ngerman", "es": "spanish", "it": "italian",
         "nl": "dutch", "pt-BR": "brazilian", "pt": "portuguese"}


def detect_language(prs):
    """babel name of the deck's main language, from its text runs (equations are ignored)."""
    counts = {}
    for slide in prs.slides:
        for run in slide._element.iter(qn("a:r")):
            rpr = run.find(qn("a:rPr"))
            lang = rpr.get("lang") if rpr is not None else None
            if lang:
                text = "".join(t.text or "" for t in run.iter(qn("a:t")))
                counts[lang] = counts.get(lang, 0) + len(text.strip())
    if not counts:
        return None
    lang = max(counts, key=counts.get)
    return BABEL.get(lang) or BABEL.get(lang.split("-")[0])


def _chart_text(obj):
    try:
        if obj.has_title and obj.axis_title.has_text_frame:
            return esc(obj.axis_title.text_frame.text.strip())
    except (AttributeError, ValueError):
        pass
    return ""


def render_chart(chart, height):
    """A native chart as a pgfplots axis, or None when the chart type is not supported."""
    try:
        kind = chart.chart_type
    except (NotImplementedError, KeyError, ValueError):
        return None
    name = str(kind).split()[0].split(".")[-1]
    if len(chart.plots) != 1:
        return None
    plot = chart.plots[0]
    lines = ["LINE", "LINE_MARKERS", "LINE_MARKERS_STACKED", "LINE_STACKED"]
    scatter = name.startswith("XY_SCATTER")
    bars = {"COLUMN_CLUSTERED": "ybar", "BAR_CLUSTERED": "xbar"}
    if name not in lines and name not in bars and not scatter:
        return None

    opts = ["width=\\linewidth", "height=%.2f\\textheight" % height]
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            opts.append("title={%s}" % esc(chart.chart_title.text_frame.text.strip()))
    except (AttributeError, ValueError):
        pass
    try:
        xlabel, ylabel = _chart_text(chart.category_axis), _chart_text(chart.value_axis)
    except (AttributeError, ValueError):
        xlabel = ylabel = ""
    if name == "BAR_CLUSTERED":
        xlabel, ylabel = ylabel, xlabel
    if xlabel:
        opts.append("xlabel={%s}" % xlabel)
    if ylabel:
        opts.append("ylabel={%s}" % ylabel)

    series = list(plot.series)
    body = []
    if scatter:
        for s in series:
            xs = [float(v.text) for v in s._element.findall(".//" + qn("c:xVal") + "//" + qn("c:v"))]
            pts = " ".join("(%g,%g)" % (x, y) for x, y in zip(xs, s.values) if y is not None)
            style = "only marks" if name == "XY_SCATTER" else "mark=none" if "NO_MARKERS" in name else ""
            body.append("\\addplot%s coordinates {%s};" % ("+[%s]" % style if style else "", pts))
    else:
        cats = [esc(str(c)).replace(",", "{,}") for c in plot.categories]
        n = len(cats)
        if name in bars:
            opts.insert(0, bars[name])
            axis = "y" if name == "BAR_CLUSTERED" else "x"
            opts += ["%stick={1,...,%d}" % (axis, n), "%sticklabels={%s}" % (axis, ", ".join(cats)),
                     "%smin=0" % ("x" if axis == "y" else "y")]  # bars start at zero
        else:
            opts += ["xtick={1,...,%d}" % n, "xticklabels={%s}" % ", ".join(cats)]
        for s in series:
            if name == "BAR_CLUSTERED":
                pts = " ".join("(%g,%d)" % (v, i) for i, v in enumerate(s.values, 1) if v is not None)
            else:
                pts = " ".join("(%d,%g)" % (i, v) for i, v in enumerate(s.values, 1) if v is not None)
            mark = "+[mark=none]" if name in ("LINE", "LINE_STACKED") else ""  # + keeps colours
            body.append("\\addplot%s coordinates {%s};" % (mark, pts))
    names = [esc(s.name or "").replace(",", "{,}") for s in series]
    if len(series) > 1 and all(names):
        # legend in one row under the plot, where it cannot hide data
        opts.append("legend style={at={(0.5,%s)}, anchor=north, legend columns=-1}"
                    % ("-0.3" if xlabel else "-0.15"))
        body.append("\\legend{%s}" % ", ".join(names))
    return ("\\begin{tikzpicture}\n\\begin{axis}[%s]\n%s\n\\end{axis}\n\\end{tikzpicture}\n"
            % (",\n  ".join(opts), "\n".join(body)))


# --------------------------------------------------------------------------
# Slide conversion
# --------------------------------------------------------------------------

class Element:
    """A piece of slide content with its position (EMU, slide coordinates).

    kind: text | table | image | drawing | comment
    ids:  ids of the top-level shapes it comes from (used to render diagrams)
    """

    def __init__(self, kind, x, y, w, h, tex=None, ids=(), movable=False):
        self.kind, self.x, self.y, self.w, self.h = kind, x, y, w, h
        self.tex = tex
        self.ids = set(ids)
        self.movable = movable     # may be absorbed into a diagram (text box / picture)
        self.src = None            # (blob, content_type, crop) for pictures
        self.render = None         # render job index for diagrams
        self.strong = False        # group or connector: always a diagram
        self.fallback = []         # what to output if the diagram cannot be rendered
        self.image = None          # (name, ok, filename) once saved
        self.anim = (None, None)   # (enter click, exit click) of a PowerPoint animation
        self.chart = None          # pgfplots code for a native chart

    @property
    def right(self):
        return self.x + self.w

    @property
    def bottom(self):
        return self.y + self.h


def image_order(elements):
    """Image elements in reading order; position k is saved as slide_<i>_image_<k+1>."""
    return [e for e in sorted(elements, key=lambda e: (e.y, e.x)) if e.kind == "image"]


def _union_box(elements):
    x0 = min(e.x for e in elements)
    y0 = min(e.y for e in elements)
    return x0, y0, max(e.right for e in elements) - x0, max(e.bottom for e in elements) - y0


class SlideConverter:
    def __init__(self, prs, images_dir, warnings, diagrams=True, fit="estimate"):
        self.prs = prs
        self.sw, self.sh = prs.slide_width, prs.slide_height
        self.images_dir = images_dir
        self.warnings = warnings
        self.diagrams = diagrams
        self.fit = fit          # "estimate": guess overfull slides; "none": leave frames alone
        self.has_charts = False
        self.anim, self.para_anim = {}, {}
        self.jobs = []  # diagrams to render: (slide index, shape ids, bbox)

    # ---- collecting elements ------------------------------------------------

    def collect(self, container, slide_no, part, transform, out, skip_ids, top_id=None):
        for shape in iter_shapes(container, self.shapes):
            if shape.shape_id in skip_ids:
                continue
            sid = top_id if top_id is not None else shape.shape_id
            nested = top_id is not None
            try:
                x, y, w, h = transform(shape.left or 0, shape.top or 0,
                                       shape.width or 0, shape.height or 0)
            except TypeError:
                x = y = w = h = 0
            stype = shape.shape_type

            def add(kind, tex=None, movable=False):
                el = Element(kind, x, y, w, h, tex, ids={sid}, movable=movable)
                el.anim = self.anim.get(sid, (None, None))
                out.append(el)
                return el

            if stype == MSO_SHAPE_TYPE.GROUP:
                children = []
                self.collect(shape._element, slide_no, part,
                             self._group_transform(shape, transform), children, skip_ids, sid)
                if nested:
                    out.extend(children)
                else:
                    el = add("drawing")
                    el.strong, el.fallback = True, children
                continue
            if shape.is_placeholder and _ph_type(shape) in SKIP_PLACEHOLDERS:
                continue

            if stype == MSO_SHAPE_TYPE.MEDIA or hasattr(shape, "poster_frame"):
                poster = getattr(shape, "poster_frame", None)
                if shape._element.find(".//" + qn("a:audioFile")) is not None:
                    add("comment", "% skipped: audio")
                elif poster is None:
                    add("comment", "% skipped: video without poster frame")
                else:
                    el = add("image", "% video poster frame")
                    el.src = (poster.blob, poster.content_type, (0, 0, 0, 0))
                continue

            if _has_image(shape):
                if shape._element.find(".//" + qn("a:audioFile")) is not None:
                    add("comment", "% skipped: audio")
                    continue
                crop = tuple(max(0.0, getattr(shape, a, 0) or 0) for a in
                             ("crop_left", "crop_top", "crop_right", "crop_bottom"))
                el = add("image", movable=not shape.is_placeholder)
                el.src = (shape.image.blob, shape.image.content_type, crop)
                continue

            if getattr(shape, "has_table", False) and shape.has_table:
                add("table", render_table(shape.table, part))
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                code = render_chart(shape.chart, min(0.75, max(0.3, h / self.sh)))
                if code:
                    add("chart").chart = code
                else:  # unsupported chart type: render it as a picture, like a diagram
                    el = add("drawing")
                    el.strong = True
                    el.fallback = [Element("comment", x, y, w, h, "% skipped: chart", ids={sid})]
                continue

            is_line = (stype == MSO_SHAPE_TYPE.LINE
                       or etree.QName(shape._element).localname == "cxnSp")
            is_autoshape = (stype in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM)
                            and not shape.is_placeholder)
            if not nested and (is_line or is_autoshape):
                el = add("drawing")
                el.strong = is_line
                # shapes this connector is glued to (e.g. the label at an input arrow)
                el.glue = {int(c.get("id")) for tag in ("a:stCxn", "a:endCxn")
                           for c in shape._element.iter(qn(tag))}
                tex = render_text_frame(shape, part) if shape.has_text_frame else ""
                if tex:
                    el.fallback = [Element("text", x, y, w, h, tex, ids={sid})]
                continue
            if is_line:
                continue

            if shape.has_text_frame:
                builds = {i: k for (spid, i), k in self.para_anim.items() if spid == shape.shape_id}
                tex = render_text_frame(shape, part, builds)
                if tex:
                    add("text", tex, movable=not shape.is_placeholder)
                continue
            if stype in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                add("comment", "% skipped: embedded object")
            elif etree.QName(shape._element).localname == "graphicFrame":
                add("comment", "% skipped: SmartArt/diagram")
                self.warnings.append(f"slide {slide_no}: SmartArt/diagram skipped")

    @staticmethod
    def _group_transform(group, parent):
        xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
        ox, oy = int(off.get("x")), int(off.get("y"))
        ex, ey = int(ext.get("cx")), int(ext.get("cy"))
        cx, cy = int(choff.get("x")), int(choff.get("y"))
        cex, cey = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
        sx, sy = ex / cex, ey / cey

        def transform(x, y, w, h):
            return parent(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy)
        return transform

    # ---- diagrams -----------------------------------------------------------

    def group_diagrams(self, elements, slide_index):
        """Merge touching drawing shapes (and the labels on top of them) into diagrams."""
        margin = 0.03 * min(self.sw, self.sh)
        drawings = [e for e in elements if e.kind == "drawing"]
        rest = [e for e in elements if e.kind != "drawing"]

        def near(a, b, m):
            return (a.x - m < b.right and b.x - m < a.right and
                    a.y - m < b.bottom and b.y - m < a.bottom)

        clusters = []
        for d in drawings:
            hits = [c for c in clusters if any(near(d, o, margin / 3) for o in c)]
            merged = [d] + [o for c in hits for o in c]
            clusters = [c for c in clusters if all(c is not h for h in hits)] + [merged]

        result = []
        for c in clusters:
            if not self.diagrams or not (len(c) >= 2 or c[0].strong):
                result.extend(f for d in c for f in d.fallback)  # keep the text only
                continue
            x, y, w, h = _union_box(c)
            box = Element("drawing", x - margin, y - margin, w + 2 * margin, h + 2 * margin)
            glued = set().union(*(getattr(d, "glue", set()) for d in c))
            absorbed = [e for e in rest if e.movable and (e.ids & glued or (
                        box.x < e.x + e.w / 2 < box.right and box.y < e.y + e.h / 2 < box.bottom))]
            rest = [e for e in rest if all(e is not a for a in absorbed)]
            members = c + absorbed
            el = Element("image", *_union_box(members), ids=set().union(*(m.ids for m in members)))
            el.fallback = [f for d in c for f in d.fallback] + absorbed
            enters = [m.anim[0] for m in members]
            if all(enters):  # every part is animated: show the picture from the first click
                el.anim = (min(enters), None)
                if len(set(enters)) > 1:
                    el.tex = "% parts animated separately in PowerPoint appear together here"
            el.render = len(self.jobs)
            self.jobs.append((slide_index, el.ids, (el.x, el.y, el.w, el.h)))
            result.append(el)
        return rest + result

    # ---- output ---------------------------------------------------------------

    def name_images(self, elements, slide_no, rendered):
        """Save pictures / rendered diagrams as slide_<i>_image_<j> in reading order."""
        final = []
        for el in elements:
            if el.render is not None and el.render not in rendered:
                final.extend(el.fallback)  # the diagram could not be rendered
            else:
                final.append(el)
        for count, el in enumerate(image_order(final), start=1):
            name = f"slide_{slide_no}_image_{count}"
            if el.render is not None:
                fname = name + ".png"
                shutil.move(rendered[el.render], os.path.join(self.images_dir, fname))
                el.image = (name, True, fname)
                continue
            fname, ok = save_image(el.src[0], el.src[1], self.images_dir, name, el.src[2])
            if not ok:
                self.warnings.append(f"slide {slide_no}: could not convert {fname} to PNG "
                                     "(install Inkscape or LibreOffice, or convert it by hand)")
            el.image = (name, ok, fname)
        return final

    def image_size(self, el, box_w, box_frac=1.0):
        """('width', fraction of \\linewidth) or ('height', fraction of \\textheight), and the
        resulting height in cm: an image keeps its share of the slide's width AND height."""
        wide = self.sw / self.sh > 1.5
        col_cm = (14.0 if wide else 10.8) * box_frac  # real width of the column
        wf = max(min(el.w / box_w, 1.0) if box_w else 1.0, 0.1)
        hf = min(0.75, 1.05 * el.h / self.sh)
        by_width_cm = wf * col_cm * el.h / max(el.w, 1)
        if by_width_cm > hf * 8.2 * 1.1:  # too tall when sized by width: size it by height
            return ("height", hf), hf * 8.2
        return ("width", wf), by_width_cm

    def image_tex(self, el, box_w, box_frac=1.0):
        name, ok, fname = el.image
        (dim, frac), _ = self.image_size(el, box_w, box_frac)
        size = ("width=%.2f\\linewidth" if dim == "width" else "height=%.2f\\textheight") % frac
        line = r"\includegraphics[%s]{%s}" % (size, name)
        if not ok:
            line = f"% TODO convert {fname} to PNG, then uncomment:\n% " + line
        return line

    def block_tex(self, elements, box_w, box_frac=1.0):
        parts = []
        for el in sorted(elements, key=lambda e: (e.y, e.x)):
            if el.kind == "image":
                img = self.image_tex(el, box_w, box_frac)
                if el.tex:
                    img = el.tex + "\n" + img
                tex = "\\begin{center}\n" + img + "\n\\end{center}"
            elif el.kind in ("table", "chart"):
                tex = "\\begin{center}\n" + el.tex + "\n\\end{center}"
            else:
                tex = el.tex
            spec = overlay_spec(*el.anim)
            if spec:  # appears or disappears on a click in PowerPoint
                tex = "\\begin{uncoverenv}%s\n%s\n\\end{uncoverenv}" % (spec, tex)
            parts.append(tex)
        return "\n\n".join(parts)

    @staticmethod
    def split_columns(elements):
        """Group elements into side-by-side columns (by horizontal overlap)."""
        cols = []
        tol = 0.02
        for el in sorted(elements, key=lambda e: e.x):
            if cols and el.x < max(e.right for e in cols[-1]) - tol * max(el.w, 1):
                cols[-1].append(el)
            else:
                cols.append([el])
        return cols

    def body_tex(self, elements):
        content = [e for e in elements if e.kind != "comment"]
        comments = [e.tex for e in elements if e.kind == "comment"]

        # a picture covering (almost) the whole slide behind other content is a background
        full = [e for e in content if e.kind == "image" and e.w * e.h > 0.85 * self.sw * self.sh]
        if full and len(content) > len(full):
            for e in full:
                comments.append(f"% background image not included: {e.image[2]}")
                content.remove(e)

        parts = []
        for kind, items in self.layout(content):
            if kind == "stack":
                parts.append(self.block_tex(items, self.sw))
                continue
            out = ["\\begin{columns}"]
            for col, span, frac in self.column_widths(items):
                out.append("\\begin{column}{%.2f\\textwidth}" % frac)
                out.append(self.block_tex(col, span, frac))
                out.append("\\end{column}")
            out.append("\\end{columns}")
            parts.append("\n".join(out))
        body = "\n\n".join(p for p in parts if p)
        return "\n".join(comments + ([body] if body else []))

    @staticmethod
    def column_widths(cols):
        """[(column, span in EMU, width as a fraction of \\textwidth)]."""
        spans = [max(e.right for e in c) - min(e.x for e in c) for c in cols]
        total = sum(spans) or 1
        return [(c, s, max(0.15, round(0.95 * s / total, 2))) for c, s in zip(cols, spans)]

    def layout(self, content):
        """[("stack", elements) | ("columns", [column elements, ...])] from top to bottom.

        Content side by side becomes columns. A full-width element above or below
        side-by-side content (e.g. the text box over a grid of plots) gets its own row.
        """
        if not content:
            return []
        cols = self.split_columns(content)
        if 2 <= len(cols) <= 4:
            return [("columns", cols)]
        wide = [e for e in content if e.w > 0.7 * self.sw]
        rest = [e for e in content if all(e is not w for w in wide)]
        cols = self.split_columns(rest) if rest else []
        if wide and 2 <= len(cols) <= 4:
            top, bottom = min(e.y for e in rest), max(e.bottom for e in rest)
            above = [e for e in wide if e.y <= top]
            below = [e for e in wide if e.y > top and e.y >= bottom - 0.05 * self.sh]
            if len(above) + len(below) == len(wide):
                rows = [("stack", above)] if above else []
                rows.append(("columns", cols))
                return rows + ([("stack", below)] if below else [])
        # a full-width text box with pictures placed in its empty right (or left) part:
        # the text really uses the rest of the width, so put them side by side
        if len(wide) == 1 and wide[0].kind == "text" and rest:
            box = wide[0]
            left, right = min(e.x for e in rest), max(e.right for e in rest)
            text = copy.copy(box)
            if left > box.x + 0.4 * box.w:
                text.w = left - box.x
                return [("columns", [[text], rest])]
            if right < box.x + 0.6 * box.w:
                text.x, text.w = right, box.right - right
                return [("columns", [rest, [text]])]
        return [("stack", content)]

    def prepare(self, slide, slide_index, is_first_title):
        """First pass: collect the slide's content; diagrams are queued for rendering."""
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = " ".join(title_shape.text_frame.text.replace("\v", " ").split())

        skip = {title_shape.shape_id} if title_shape is not None else set()
        subtitle = ""
        if is_first_title:
            for ph in slide.placeholders:
                if _ph_type(ph) == PP_PLACEHOLDER.SUBTITLE:
                    subtitle = " ".join(ph.text_frame.text.replace("\v", " ").split())
                    skip.add(ph.shape_id)

        elements = []
        identity = lambda x, y, w, h: (x, y, w, h)
        self.shapes = slide.shapes
        self.anim, self.para_anim = read_animations(slide)
        self.collect(slide.shapes._spTree, slide_index + 1, slide.part, identity, elements, skip)
        if is_first_title:  # logos etc. on the title slide are only extracted
            elements = [e for e in elements if e.kind == "image"]
        else:
            elements = self.group_diagrams(elements, slide_index)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        # when something disappears on the last click, Beamer must be told to make that page
        clicks = [k for pair in self.anim.values() for k in pair if k] + list(self.para_anim.values())
        exits = any(leave for _, leave in self.anim.values())
        return dict(title=title, subtitle=subtitle, elements=elements, notes=notes,
                    first_title=is_first_title, overlays=max(clicks) if exits and clicks else None)

    def write_charts(self, elements, slide_no):
        """Save native charts as charts/slide_<i>_chart_<j>.tex (pgfplots)."""
        charts = [e for e in sorted(elements, key=lambda e: (e.y, e.x)) if e.kind == "chart"]
        if not charts:
            return
        cdir = os.path.join(os.path.dirname(self.images_dir), "charts")
        os.makedirs(cdir, exist_ok=True)
        for j, el in enumerate(charts, start=1):
            name = f"slide_{slide_no}_chart_{j}"
            with open(os.path.join(cdir, name + ".tex"), "w", encoding="utf-8") as f:
                f.write(el.chart)
            el.tex = "\\input{charts/%s}" % name
        self.has_charts = True

    # ---- overfull slides --------------------------------------------------------

    LINES = 14.5  # lines of normal text that fit below the frame title

    def _lines(self, el, width, frac=1.0):
        """Rough height of an element, in lines of text, when set `width` EMU wide."""
        wide = self.sw / self.sh > 1.5
        col_cm = (14.0 if wide else 10.8) * frac
        if el.kind in ("text", "comment"):
            total = 0.0
            per_line = max(10, col_cm * 5.1)  # characters per line at ~11pt
            for line in (el.tex or "").splitlines():
                s = line.strip()
                if not s:
                    total += 0.4
                elif s.startswith("\\["):
                    total += 3.0 if "\\frac" in s else 2.2
                elif s.startswith(("\\begin", "\\end", "%")):
                    total += 0.2
                else:
                    visible = re.sub(r"\\[a-zA-Z]+|[{}$]|<\d*-?\d*>", "", s)
                    total += max(1, math.ceil(len(visible) / per_line))
            return total
        if el.kind == "table":
            return el.tex.count("\\hline") * 1.1
        if el.kind == "chart":
            return min(0.75, el.h / self.sh) * 8.6 / 0.48
        if el.kind == "image":
            return self.image_size(el, width, frac)[1] / 0.48 + 0.5
        return 0

    def estimate_fit(self, elements):
        """Frame options for a slide that probably does not fit, e.g. 'shrink=20'."""
        content = [e for e in elements if e.kind != "comment"]
        if not content:
            return ""
        need = 0.0
        rows = self.layout(content)
        for kind, items in rows:
            if kind == "stack":
                need += sum(self._lines(e, self.sw) for e in items)
            else:
                need += max(sum(self._lines(e, span, frac) for e in c)
                            for c, span, frac in self.column_widths(items))
        if need <= 1.25 * self.LINES:
            return ""
        if rows == [("stack", content)] and all(e.kind == "text" for e in content) and \
                "\\[" not in "".join(e.tex for e in content):
            return "allowframebreaks"
        # same formula as for a measured overflow; ~6.5pt per line beyond ~16.5 lines
        # (calibrated against LaTeX's own measurements on real lecture decks)
        too_high = (need - 16.5) * 6.5
        return "shrink=%d" % min(40, math.ceil(100 * too_high / (too_high + 200)) + 3)

    def finish(self, state, slide_no, rendered):
        """Second pass: save the images and write the frame."""
        elements = self.name_images(state["elements"], slide_no, rendered)
        self.write_charts(elements, slide_no)
        lines = []
        if state["first_title"]:
            lines.append("\\begin{frame}")
            lines.append("\\titlepage")
            for e in elements:
                lines.append(f"% not included from title slide: {e.image[2]}")
        else:
            title = state["title"]
            fit = self.estimate_fit(elements) if self.fit == "estimate" else ""
            opts = "[%s]" % fit if fit else ""
            if state.get("overlays"):
                opts = "<1-%d>" % state["overlays"] + opts
            lines.append("\\begin{frame}%s{%s}" % (opts, esc(title)) if title else "\\begin{frame}" + opts)
            body = self.body_tex(elements)
            if body:
                lines.append(body)
        if state["notes"]:
            lines.append("\\note{%s}" % esc(state["notes"]).replace("\n", "\n\n"))
        lines.append("\\end{frame}")
        return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------
# Rendering diagrams with PowerPoint (Windows)
# --------------------------------------------------------------------------

RENDER_PS1 = r"""
param([string]$pptx, [string]$outdir, [string]$slides, [int]$width, [int]$height)
$ErrorActionPreference = "Stop"
$app = New-Object -ComObject PowerPoint.Application
$pres = $null
try {
    $pres = $app.Presentations.Open($pptx, -1, 0, 0)
    foreach ($n in $slides.Split(",")) {
        $pres.Slides.Item([int]$n).Export((Join-Path $outdir "s$n.png"), "PNG", $width, $height)
    }
} finally {
    if ($pres) { $pres.Close() }
    if ($app.Presentations.Count -eq 0) { $app.Quit() }
}
"""


def _element_id(el):
    cnvpr = el.find(".//" + qn("p:cNvPr"))
    return int(cnvpr.get("id")) if cnvpr is not None else None


def render_diagrams(pptx_path, jobs, background, tmpdir, warnings):
    """Render each job (slide index, shape ids, bbox) to a cropped PNG using PowerPoint.

    Works on a temporary copy of the deck in which every slide keeps only its diagram
    shapes, so nothing else ends up in the picture. Returns {job index: png path}.
    """
    if not jobs:
        return {}
    if os.name != "nt" or not shutil.which("powershell") or Image is None:
        warnings.append("diagrams kept as text: rendering them needs PowerPoint on Windows")
        return {}

    prs = Presentation(pptx_path)
    keep = {}
    for slide_index, ids, _ in jobs:
        keep.setdefault(slide_index, set()).update(ids)
    bg_xml = ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr>'
              '<a:solidFill><a:srgbClr val="%02X%02X%02X"/></a:solidFill><a:effectLst/>'
              '</p:bgPr></p:bg>' % tuple(background))
    for index, slide in enumerate(prs.slides):
        sld = slide._element
        sld.attrib.pop("show", None)
        sld.set("showMasterSp", "0")
        for timing in sld.findall(qn("p:timing")):  # animations would point at removed shapes
            sld.remove(timing)
        tree = slide.shapes._spTree
        for child in list(tree):
            if etree.QName(child).localname in ("nvGrpSpPr", "grpSpPr", "extLst"):
                continue
            if _element_id(child) not in keep.get(index, ()):
                tree.remove(child)
        csld = sld.find(qn("p:cSld"))
        old = csld.find(qn("p:bg"))
        if old is not None:
            csld.remove(old)
        csld.insert(0, etree.fromstring(bg_xml))
    deck = os.path.join(tmpdir, "render.pptx")
    prs.save(deck)

    width = 2400
    height = round(width * prs.slide_height / prs.slide_width)
    script = os.path.join(tmpdir, "render.ps1")
    with open(script, "w", encoding="utf-8") as f:
        f.write(RENDER_PS1)
    slides = ",".join(str(i + 1) for i in sorted(keep))
    print(f"Rendering {len(jobs)} diagram(s) with PowerPoint...")
    try:
        proc = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                               "-File", script, "-pptx", os.path.abspath(deck),
                               "-outdir", os.path.abspath(tmpdir), "-slides", slides,
                               "-width", str(width), "-height", str(height)],
                              capture_output=True, text=True, timeout=600)
    except (OSError, subprocess.SubprocessError) as e:
        warnings.append(f"diagrams kept as text: PowerPoint rendering failed ({e})")
        return {}
    if proc.returncode != 0:
        msg = (proc.stderr or proc.stdout).strip().splitlines()
        warnings.append("diagrams kept as text: PowerPoint rendering failed (%s)"
                        % (msg[0] if msg else "unknown error"))
        return {}

    out = {}
    scale = width / prs.slide_width
    pad = int(0.05 * width)
    per_slide = {}
    for slide_index, _, _ in jobs:
        per_slide[slide_index] = per_slide.get(slide_index, 0) + 1
    for key, (slide_index, _, (x, y, w, h)) in enumerate(jobs):
        png = os.path.join(tmpdir, f"s{slide_index + 1}.png")
        if not os.path.exists(png):
            continue
        with Image.open(png) as im:
            im = im.convert("RGB")
            if per_slide[slide_index] == 1:
                # only this diagram is on the rendered slide: trim below finds its extent
                # (connectors may be drawn outside their bounding box)
                crop = im.copy()
            else:
                crop = im.crop((max(0, int(x * scale) - pad), max(0, int(y * scale) - pad),
                                min(im.width, int((x + w) * scale) + pad),
                                min(im.height, int((y + h) * scale) + pad)))
        # trim to the drawn content, keeping a small border
        diff = ImageChops.difference(crop, Image.new("RGB", crop.size, tuple(background)))
        content = diff.convert("L").point(lambda v: 255 if v > 12 else 0).getbbox()
        if content:
            b = int(0.01 * width)
            crop = crop.crop((max(0, content[0] - b), max(0, content[1] - b),
                              min(crop.width, content[2] + b), min(crop.height, content[3] + b)))
        path = os.path.join(tmpdir, f"job{key}.png")
        crop.save(path)
        out[key] = path
    return out


SHAPE_TAGS = {"sp", "pic", "grpSp", "graphicFrame", "cxnSp"}


def iter_shapes(container, parent):
    """Shapes in a shape tree or group, including those wrapped in mc:AlternateContent
    (PowerPoint uses it for text boxes containing equations; python-pptx skips them)."""
    for child in container:
        tag = etree.QName(child).localname
        if tag == "AlternateContent":
            choice = child.find(MC + "Choice")
            if choice is not None:
                yield from iter_shapes(choice, parent)
        elif tag in SHAPE_TAGS:
            yield SlideShapeFactory(child, parent)


def _has_image(shape):
    if shape.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return False
    try:
        return shape.image is not None
    except (AttributeError, ValueError, KeyError):
        return False


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def build_main(prs, theme, color_line, n_slides, hidden, title, subtitle,
               sections=None, language=None, title_slide=None, charts=False, outline=True):
    ratio = prs.slide_width / prs.slide_height
    aspect = "169" if abs(ratio - 16 / 9) < 0.05 else "1610" if abs(ratio - 1.6) < 0.03 else "43"
    author = prs.core_properties.author or ""
    sections = sections or {}
    lines = [
        "\\documentclass[aspectratio=%s]{beamer}" % aspect,
        "\\usetheme{%s}" % theme,
    ]
    lines += color_line
    if language and language != "english":
        # T1 fonts are for pdflatex only (XeLaTeX/LuaLaTeX, e.g. Tectonic, handle Unicode already)
        lines += ["\\usepackage{iftex}",
                  "\\ifPDFTeX\\usepackage[T1]{fontenc}\\usepackage{lmodern}\\fi",
                  "\\usepackage[%s]{babel}" % language]
    if charts:
        lines += ["\\usepackage{pgfplots}", "\\pgfplotsset{compat=1.18}"]
        if language and language != "english":
            lines.append("\\usetikzlibrary{babel}")
    lines += [
        "\\graphicspath{{images/}}",
        "",
        "\\title{%s}" % esc(title),
    ]
    if subtitle:
        lines.append("\\subtitle{%s}" % esc(subtitle))
    lines += ["\\author{%s}" % esc(author), "", "\\begin{document}", ""]
    outline_frame = "\\begin{frame}{%s}\\tableofcontents\\end{frame}" % (
        "Plan" if language == "french" else "Outline")
    if sections and outline and title_slide is None:
        lines.append(outline_frame)
    for i in range(1, n_slides + 1):
        if i - 1 in sections:
            lines.append("")
            lines.append("\\section{%s}" % esc(sections[i - 1]))
        prefix = "% (hidden) " if i in hidden else ""
        lines.append(prefix + "\\input{slides/slide_%d}" % i)
        if sections and outline and title_slide == i - 1:
            lines.append(outline_frame)
    lines += ["", "\\end{document}", ""]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# Compiling
# --------------------------------------------------------------------------

def find_engine():
    """(name, path) of latexmk, pdflatex or tectonic; (None, None) if there is none."""
    for name in ("latexmk", "pdflatex", "tectonic"):
        path = shutil.which(name)
        if path:
            return name, path
    local = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "tectonic", "tectonic.exe")
    if os.path.exists(local):
        return "tectonic", local
    return None, None


def compile_project(out_dir):
    """Compile main.tex. Returns dict(engine, pdf, errors [(file, line, message)],
    overfull {slide number: points too high}), or None when no LaTeX engine is installed."""
    name, exe = find_engine()
    if exe is None:
        return None
    flags = ["-interaction=nonstopmode", "-file-line-error"]
    if name == "latexmk":
        runs = [[exe, "-pdf"] + flags + ["main.tex"]]
    elif name == "pdflatex":
        runs = [[exe] + flags + ["main.tex"]] * 2  # twice, for the outline and page numbers
    else:
        runs = [[exe, "--keep-logs", "main.tex"]]
    output = ""
    for cmd in runs:
        try:
            proc = subprocess.run(cmd, cwd=out_dir, capture_output=True, text=True,
                                  encoding="utf-8", errors="replace", timeout=1800)
            output += proc.stdout + proc.stderr
        except (OSError, subprocess.SubprocessError) as e:
            output += str(e)
    log_path = os.path.join(out_dir, "main.log")
    log = open(log_path, encoding="utf-8", errors="replace").read() if os.path.exists(log_path) else ""

    # walk the log: remember which slide file was opened last, attribute messages to it
    errors, overfull, wide = [], {}, []
    current, pending = None, None
    token = re.compile(r"\((?:\./)?((?:slides|diagrams|charts)/[\w-]+)(?:\.tex)?"
                       r"|^(?:\./)?((?:slides|diagrams|charts)/[\w-]+\.tex|main\.tex):(\d+): (.*)$"
                       r"|^! (.*)$|^l\.(\d+)|Overfull \\vbox \(([\d.]+)pt too high\)"
                       r"|Overfull \\hbox \(([\d.]+)pt too wide\) detected at line (\d+)", re.M)
    for m in token.finditer(log):
        if m.group(1):
            current = m.group(1) + ".tex"
        elif m.group(2):
            errors.append((m.group(2), int(m.group(3)), m.group(4).strip()))
            pending = None
        elif m.group(5):
            pending = m.group(5).strip()
        elif m.group(6) and pending:
            errors.append((current or "main.tex", int(m.group(6)), pending))
            pending = None
        elif m.group(7) and current:
            slide = re.search(r"slide_(\d+)", current)
            if slide:
                n = int(slide.group(1))
                overfull[n] = max(overfull.get(n, 0), float(m.group(7)))
        elif m.group(8) and current and float(m.group(8)) > 10:
            wide.append((current, int(m.group(9)), float(m.group(8))))
    for m in re.finditer(r"^error: (?:\./)?([\w/.-]+\.tex):(\d+): (.*)$", output, re.M):
        errors.append((m.group(1), int(m.group(2)), m.group(3).strip()))
    unique = []
    for e in errors:
        if e not in unique:
            unique.append(e)
    pdf = os.path.join(out_dir, "main.pdf")
    return dict(engine=name, pdf=pdf if os.path.exists(pdf) else None, errors=unique,
                overfull=overfull, wide=wide, output=output)


def fit_equation(out_dir, fname, line):
    """Scale a display equation that is wider than the slide down to the line width."""
    path = os.path.join(out_dir, fname)
    if not os.path.exists(path):
        return False
    lines = open(path, encoding="utf-8").read().split("\n")
    if not 0 < line <= len(lines):
        return False
    equation = re.compile(r"^(\s*)\\\[ (.*) \\\]\s*$")
    # LaTeX reports the equation's own line, or the end of the column that holds it:
    # then take the widest equation between that line and the start of the column
    candidates = []
    for i in range(line - 1, -1, -1):
        m = equation.match(lines[i])
        if m and "\\resizebox" not in m.group(2):
            candidates.append((len(re.sub(r"\\[a-zA-Z]+|[{}]", "", m.group(2))), i, m))
        if i < line - 1 and re.match(r"\s*\\begin\{(column|frame)\}", lines[i]):
            break
        if i == line - 1 and candidates:
            break
    if not candidates:
        return False
    _, i, m = max(candidates, key=lambda c: c[0])
    lines[i] = "%s\\[ \\resizebox{\\linewidth}{!}{$\\displaystyle %s$} \\]" % m.groups()
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return True


def add_shrink(slides_dir, slide_no, too_high):
    """Shrink a frame that is `too_high` points taller than the page allows."""
    path = os.path.join(slides_dir, f"slide_{slide_no}.tex")
    if not os.path.exists(path):
        return None
    text = open(path, encoding="utf-8").read()
    m = re.match(r"\\begin\{frame\}(<[^>]*>)?(\[([^\]]*)\])?", text)
    if not m:
        return None
    overlays = m.group(1) or ""
    opts = [o.strip() for o in (m.group(3) or "").split(",") if o.strip()]
    if "allowframebreaks" in opts:
        return None
    old = next((int(o.split("=")[1]) for o in opts if o.startswith("shrink=")), 0)
    new = min(60, old + math.ceil(100 * too_high / (too_high + 200)) + 3)
    text_only = not re.search(r"\\includegraphics|\\input|\\\[|tabular", text)
    if new > 30 and text_only:  # a long text slide reads better split over two frames
        opts = [o for o in opts if not o.startswith("shrink")] + ["allowframebreaks"]
        new = "split"
    else:
        opts = [o for o in opts if not o.startswith("shrink")] + ["shrink=%d" % new]
    text = "\\begin{frame}%s[%s]" % (overlays, ", ".join(opts)) + text[m.end():]
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return new


def compile_and_fit(out_dir):
    """Compile, shrink the frames LaTeX reports as too tall, compile again, and report."""
    result = compile_project(out_dir)
    if result is None:
        print("--compile: no LaTeX engine found. Install one of: TeX Live or MiKTeX (pdflatex), "
              "or Tectonic (https://tectonic-typesetting.github.io).")
        return 1
    shrunk, scaled = {}, set()
    for _ in range(3):
        if result["errors"]:
            break
        changed = False
        for fname, line, _pt in result["wide"]:  # equations wider than the slide
            if fit_equation(out_dir, fname, line):
                scaled.add(fname)
                changed = True
        for slide, pt in sorted(result["overfull"].items()):  # frames taller than the slide
            value = add_shrink(os.path.join(out_dir, "slides"), slide, pt)
            if value:
                shrunk[slide] = value
                changed = True
        if not changed:
            break
        result = compile_project(out_dir)
    print(f"Compiled with {result['engine']}: " + (result["pdf"] or "no PDF produced"))
    if shrunk:
        print("  made to fit: " + ", ".join(
            f"slide {s} ({'split over frames' if v == 'split' else f'shrunk {v}%'})"
            for s, v in sorted(shrunk.items())))
    if scaled:
        print("  wide equations scaled down in: " + ", ".join(sorted(scaled)))
    if result["overfull"] and not result["errors"]:
        print("  still too full: " + ", ".join(f"slide {s}" for s in sorted(result["overfull"])))
    for f, line, msg in result["errors"]:
        print(f"  error: {f}:{line}: {msg}")
    if result["pdf"] is None and not result["errors"]:
        print(result["output"][-2000:])
    return 0 if result["pdf"] and not result["errors"] else 1


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convert a PowerPoint (.pptx) file to a Beamer project.")
    ap.add_argument("pptx", nargs="?", help="input .pptx file")
    ap.add_argument("-o", "--output", help="output folder (default: <name>_beamer)")
    ap.add_argument("--theme", default="Madrid", help="Beamer theme, e.g. Madrid, Warsaw, metropolis")
    ap.add_argument("--color", default="auto",
                    help="colour theme: auto (closest to the deck), exact (deck colour), "
                         "none, or a Beamer colour theme name such as beaver")
    ap.add_argument("--sections", default="auto", choices=["auto", "pptx", "none"],
                    help="where \\section comes from: auto (PowerPoint sections, section-header "
                         "slides or numbered titles like '2.1 ...'), pptx (PowerPoint sections only), none")
    ap.add_argument("--no-outline", action="store_true", help="no outline slide after the title")
    ap.add_argument("--language", default="auto",
                    help="babel language: auto (from the slides), none, or a name such as french")
    ap.add_argument("--compile", action="store_true",
                    help="compile the result (latexmk, pdflatex or tectonic), report errors by "
                         "slide file and shrink slides that are too full")
    ap.add_argument("--no-render", action="store_true",
                    help="do not render drawn diagrams with PowerPoint; keep their text only")
    ap.add_argument("--list-themes", action="store_true", help="list known themes and exit")
    args = ap.parse_args(argv)

    if args.list_themes:
        print("Themes:       " + ", ".join(THEMES))
        print("Colour themes: " + ", ".join(list(COLOR_THEMES) + OTHER_COLOR_THEMES))
        return 0
    if not args.pptx:
        ap.error("the input .pptx file is required")

    prs = Presentation(args.pptx)
    out_dir = args.output or os.path.splitext(args.pptx)[0] + "_beamer"
    slides_dir = os.path.join(out_dir, "slides")
    images_dir = os.path.join(out_dir, "images")
    os.makedirs(slides_dir, exist_ok=True)
    os.makedirs(images_dir, exist_ok=True)

    # colours
    accent, background = deck_colors(prs)
    exact = "\\usecolortheme[RGB={%d,%d,%d}]{structure}" % accent
    if args.color == "exact":
        color_line = [exact]
    elif args.color == "none":
        color_line = []
    else:
        name = pick_color_theme(accent, background) if args.color == "auto" else args.color
        color_line = ["\\usecolortheme{%s}" % name,
                      "%% exact deck colour instead: %s" % exact]

    language = {"auto": detect_language(prs), "none": None}.get(args.language, args.language)
    sections = find_sections(prs, args.sections)

    warnings = []
    conv = SlideConverter(prs, images_dir, warnings, diagrams=not args.no_render,
                          fit="none" if args.compile else "estimate")
    title = prs.core_properties.title or ""
    if title.lower() in ("", "powerpoint presentation", "présentation powerpoint"):
        title = os.path.splitext(os.path.basename(args.pptx))[0]  # PowerPoint's default title
    subtitle = ""
    hidden = set()
    title_slide = None

    # pass 1: read every slide, queue diagrams for rendering
    states = []
    for index, slide in enumerate(prs.slides):
        if slide._element.get("show") in ("0", "false"):
            hidden.add(index + 1)
        first_title = title_slide is None and _is_title_slide(slide)
        states.append(conv.prepare(slide, index, first_title))
        if first_title:
            title_slide = index
            title = states[-1]["title"] or title
            subtitle = states[-1]["subtitle"]

    # render diagrams, then pass 2: write the slides
    with tempfile.TemporaryDirectory() as tmp:
        rendered = render_diagrams(args.pptx, conv.jobs, background, tmp, warnings)
        for n, state in enumerate(states, start=1):
            with open(os.path.join(slides_dir, f"slide_{n}.tex"), "w", encoding="utf-8") as f:
                f.write(conv.finish(state, n, rendered))
    n = len(states)

    with open(os.path.join(out_dir, "main.tex"), "w", encoding="utf-8") as f:
        f.write(build_main(prs, args.theme, color_line, n, hidden, title, subtitle, sections,
                           language, title_slide, conv.has_charts, not args.no_outline))

    print(f"Converted {n} slides -> {out_dir}")
    print("Deck accent colour: RGB%s, background: RGB%s" % (accent, background))
    if color_line:
        print("Colour: " + color_line[0])
    if language:
        print("Language: " + language)
    if sections:
        print(f"Sections: {len(sections)}")
    for w in warnings:
        print("warning: " + w, file=sys.stderr)
    if args.compile:
        return compile_and_fit(out_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())

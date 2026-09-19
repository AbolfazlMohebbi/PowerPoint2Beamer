#!/usr/bin/env python3
"""Convert a PowerPoint (.pptx) presentation into a simple Beamer LaTeX project.

Output:
    OUTDIR/main.tex                     preamble + \\input of every slide
    OUTDIR/slides/slide_<i>.tex         one frame per slide
    OUTDIR/images/slide_<i>_image_<j>.* pictures and video poster frames

Usage:
    python pptx2beamer.py deck.pptx [-o OUTDIR] [--theme Madrid] [--color auto]
"""

import argparse
import io
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.shapetree import SlideShapeFactory

try:
    from PIL import Image, ImageChops
except ImportError:  # Pillow is optional; only needed for format conversion/cropping
    Image = None


# --------------------------------------------------------------------------
# Themes
# --------------------------------------------------------------------------

THEMES = [
    "default", "AnnArbor", "Antibes", "Bergen", "Berkeley", "Berlin", "Boadilla",
    "CambridgeUS", "Copenhagen", "Darmstadt", "Dresden", "Frankfurt", "Goettingen",
    "Hannover", "Ilmenau", "JuanLesPins", "Luebeck", "Madrid", "Malmoe", "Marburg",
    "Montpellier", "PaloAlto", "Pittsburgh", "Rochester", "Singapore", "Szeged",
    "Warsaw", "metropolis",
]

# Beamer colour themes with the (approximate) dominant colour they paint the
# headline/frame title with, and their page background colour.
COLOR_THEMES = {
    "whale":     ((51, 51, 179),   (255, 255, 255)),
    "dolphin":   ((140, 140, 215), (255, 255, 255)),
    "seahorse":  ((200, 200, 235), (255, 255, 255)),
    "beaver":    ((170, 0, 0),     (255, 255, 255)),
    "crane":     ((255, 191, 0),   (255, 255, 255)),
    "wolverine": ((255, 204, 51),  (255, 255, 255)),
    "spruce":    ((0, 90, 50),     (255, 255, 255)),
    "seagull":   ((150, 150, 150), (255, 255, 255)),
    "dove":      ((60, 60, 60),    (255, 255, 255)),
    "beetle":    ((51, 51, 153),   (153, 153, 153)),
    "fly":       ((220, 220, 220), (77, 77, 77)),
    "albatross": ((150, 150, 255), (20, 20, 60)),
}
OTHER_COLOR_THEMES = ["default", "orchid", "rose", "lily", "structure", "sidebartab"]


# --------------------------------------------------------------------------
# LaTeX escaping
# --------------------------------------------------------------------------

LATEX_ESCAPES = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
    "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}", "<": r"\textless{}", ">": r"\textgreater{}",
    # common Unicode symbols pdflatex does not know out of the box
    "\u2022": r"\textbullet{}", "\u2013": "--", "\u2014": "---", "\u2026": r"\ldots{}",
    "\u2192": r"$\rightarrow$", "\u2190": r"$\leftarrow$", "\u21d2": r"$\Rightarrow$",
    "\u2264": r"$\leq$", "\u2265": r"$\geq$", "\u2260": r"$\neq$", "\u00d7": r"$\times$",
    "\u00b1": r"$\pm$", "\u2248": r"$\approx$", "\u221e": r"$\infty$", "\u2211": r"$\sum$",
    "\u03b1": r"$\alpha$", "\u03b2": r"$\beta$", "\u03b3": r"$\gamma$", "\u03b4": r"$\delta$",
    "\u03b5": r"$\epsilon$", "\u03b8": r"$\theta$", "\u03bb": r"$\lambda$", "\u03bc": r"$\mu$",
    "\u03c0": r"$\pi$", "\u03c3": r"$\sigma$", "\u03c6": r"$\phi$", "\u03c9": r"$\omega$",
    "\u0394": r"$\Delta$", "\u03a3": r"$\Sigma$", "\u03a9": r"$\Omega$",
    "\u2713": r"$\checkmark$", "\u00a0": "~", "\u200b": "",
}


def esc(text):
    out = []
    for ch in text:
        if 0xF000 <= ord(ch) <= 0xF0FF:  # Symbol/Wingdings private-use glyphs
            continue
        if 0x1D400 <= ord(ch) <= 0x1D7FF or ch in "ℎℒ":  # math letters (𝑠, ℒ, ...)
            out.append("$" + _math_char(ch).strip() + "$")
            continue
        out.append(LATEX_ESCAPES.get(ch, ch))
    return "".join(out)


def esc_url(url):
    return url.replace("\\", "/").replace("%", r"\%").replace("#", r"\#")


# --------------------------------------------------------------------------
# Colours
# --------------------------------------------------------------------------

def _lab(rgb):
    def lin(c):
        c /= 255.0
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (lin(c) for c in rgb)
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883

    def f(t):
        return t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    fx, fy, fz = f(x), f(y), f(z)
    return 116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz)


def color_distance(a, b):
    return math.dist(_lab(a), _lab(b))


def hex_rgb(val):
    return tuple(int(val[i:i + 2], 16) for i in (0, 2, 4))


def deck_colors(prs):
    """Return (accent, background) RGB tuples of the deck's theme."""
    master = prs.slide_master
    scheme = {}
    try:
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        for el in theme.find(".//" + qn("a:clrScheme")):
            name = etree.QName(el).localname
            srgb = el.find(qn("a:srgbClr"))
            sysc = el.find(qn("a:sysClr"))
            if srgb is not None:
                scheme[name] = hex_rgb(srgb.get("val"))
            elif sysc is not None and sysc.get("lastClr"):
                scheme[name] = hex_rgb(sysc.get("lastClr"))
    except (KeyError, TypeError, AttributeError):
        pass

    clr_map = master._element.find(qn("p:clrMap"))
    aliases = dict(clr_map.attrib) if clr_map is not None else {}
    aliases.setdefault("bg1", "lt1")
    aliases.setdefault("tx1", "dk1")
    aliases.setdefault("bg2", "lt2")
    aliases.setdefault("tx2", "dk2")

    def resolve(fill_parent):
        if fill_parent is None:
            return None
        srgb = fill_parent.find(".//" + qn("a:srgbClr"))
        if srgb is not None:
            return hex_rgb(srgb.get("val"))
        sch = fill_parent.find(".//" + qn("a:schemeClr"))
        if sch is not None:
            name = sch.get("val")
            return scheme.get(aliases.get(name, name))
        return None

    # background: first slide that is not a title slide, then its layout, then master
    background = None
    sample = next((s for s in prs.slides if not _is_title_slide(s)), None)
    candidates = [sample, sample.slide_layout] if sample is not None else []
    for obj in candidates + [master]:
        bg = obj._element.find(qn("p:cSld")).find(qn("p:bg"))
        if bg is not None:
            background = resolve(bg)
            if background:
                break
    background = background or scheme.get("lt1", (255, 255, 255))
    accent = scheme.get("accent1") or scheme.get("dk2") or (51, 51, 179)
    return accent, background


def pick_color_theme(accent, background):
    def score(item):
        struct, bg = item[1]
        return color_distance(accent, struct) + color_distance(background, bg)
    return min(COLOR_THEMES.items(), key=score)[0]


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------

CONTENT_TYPE_EXT = {
    "image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif",
    "image/bmp": "bmp", "image/tiff": "tiff", "image/x-emf": "emf", "image/emf": "emf",
    "image/x-wmf": "wmf", "image/wmf": "wmf", "image/svg+xml": "svg",
    "application/pdf": "pdf",
}


def _convert_vector(blob, ext, out_png):
    """Convert EMF/WMF to PNG. Returns True on success."""
    if Image is not None:
        try:
            im = Image.open(io.BytesIO(blob))
            im.load(dpi=200)
            im.save(out_png)
            return True
        except Exception:
            pass
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in." + ext)
        with open(src, "wb") as f:
            f.write(blob)
        cmds = []
        if shutil.which("inkscape"):
            cmds.append(["inkscape", src, "--export-type=png",
                         "--export-filename=" + os.path.join(tmp, "in.png")])
        soffice = shutil.which("soffice") or shutil.which("soffice.exe")
        if soffice:
            cmds.append([soffice, "--headless", "--convert-to", "png", "--outdir", tmp, src])
        for cmd in cmds:
            try:
                subprocess.run(cmd, capture_output=True, timeout=120)
            except (OSError, subprocess.SubprocessError):
                continue
            produced = os.path.join(tmp, "in.png")
            if os.path.exists(produced):
                shutil.move(produced, out_png)
                return True
    return False


def save_image(blob, content_type, images_dir, name, crop=(0, 0, 0, 0)):
    """Save an image blob as images_dir/name.<ext> in a pdflatex-friendly format.

    Returns the file name that was written, and whether pdflatex can include it.
    """
    ext = CONTENT_TYPE_EXT.get(content_type, "png")
    cropped = any(c > 0 for c in crop)

    if ext in ("png", "jpg", "pdf") and not cropped:
        fname = f"{name}.{ext}"
        with open(os.path.join(images_dir, fname), "wb") as f:
            f.write(blob)
        return fname, True

    if ext in ("emf", "wmf"):
        fname = f"{name}.png"
        if _convert_vector(blob, ext, os.path.join(images_dir, fname)):
            return fname, True
        fname = f"{name}.{ext}"
        with open(os.path.join(images_dir, fname), "wb") as f:
            f.write(blob)
        return fname, False

    if Image is not None and ext != "svg":
        try:
            im = Image.open(io.BytesIO(blob))
            if cropped:
                w, h = im.size
                l, t, r, b = crop
                im = im.crop((int(l * w), int(t * h), int(w - r * w), int(h - b * h)))
            out_ext = "jpg" if ext == "jpg" else "png"
            if out_ext == "png" and im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
                im = im.convert("RGBA")
            fname = f"{name}.{out_ext}"
            im.save(os.path.join(images_dir, fname))
            return fname, True
        except Exception:
            pass

    fname = f"{name}.{ext}"
    with open(os.path.join(images_dir, fname), "wb") as f:
        f.write(blob)
    return fname, ext in ("png", "jpg", "pdf")


# --------------------------------------------------------------------------
# Office equations (OMML) -> LaTeX math
# --------------------------------------------------------------------------

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

MATH_SYMBOLS = {
    "−": "-", "∗": "*", "×": r"\times ", "·": r"\cdot ", "⋅": r"\cdot ",
    "÷": r"\div ", "±": r"\pm ", "∓": r"\mp ", "∞": r"\infty ",
    "→": r"\rightarrow ", "←": r"\leftarrow ", "⇒": r"\Rightarrow ",
    "⇔": r"\Leftrightarrow ", "↔": r"\leftrightarrow ", "↦": r"\mapsto ",
    "≤": r"\leq ", "≥": r"\geq ", "≠": r"\neq ", "≈": r"\approx ",
    "≡": r"\equiv ", "∝": r"\propto ", "∼": r"\sim ", "≪": r"\ll ", "≫": r"\gg ",
    "∈": r"\in ", "∉": r"\notin ", "⊂": r"\subset ", "⊆": r"\subseteq ",
    "∪": r"\cup ", "∩": r"\cap ", "∀": r"\forall ", "∃": r"\exists ",
    "∂": r"\partial ", "∇": r"\nabla ", "…": r"\ldots ", "⋯": r"\cdots ",
    "⋮": r"\vdots ", "⋱": r"\ddots ", "′": "'", "″": "''", "°": r"^\circ ",
    "∘": r"\circ ", "∑": r"\sum ", "∏": r"\prod ", "∫": r"\int ",
    "√": r"\sqrt ", "‖": r"\| ", "⟨": r"\langle ", "⟩": r"\rangle ",
    "ℏ": r"\hbar ", "ℓ": r"\ell ", "∅": r"\emptyset ", "∧": r"\wedge ",
    "∨": r"\vee ", "¬": r"\neg ", "∥": r"\parallel ", "⊥": r"\perp ",
    " ": " ", " ": r"\, ", "​": "", "⁡": "", "⁢": "", "⁣": "",
    "{": r"\{", "}": r"\}", "%": r"\%", "#": r"\#", "$": r"\$", "_": r"\_", "~": r"\sim ",
    "\\": r"\backslash ",
}
GREEK = {
    "α": "alpha", "β": "beta", "γ": "gamma", "δ": "delta", "ε": "epsilon",
    "ϵ": "epsilon", "ζ": "zeta", "η": "eta", "θ": "theta", "ϑ": "vartheta",
    "ι": "iota", "κ": "kappa", "λ": "lambda", "μ": "mu", "ν": "nu",
    "ξ": "xi", "π": "pi", "ρ": "rho", "σ": "sigma", "ς": "varsigma",
    "τ": "tau", "υ": "upsilon", "φ": "phi", "ϕ": "phi", "χ": "chi",
    "ψ": "psi", "ω": "omega", "Γ": "Gamma", "Δ": "Delta", "Θ": "Theta",
    "Λ": "Lambda", "Ξ": "Xi", "Π": "Pi", "Σ": "Sigma", "Υ": "Upsilon",
    "Φ": "Phi", "Ψ": "Psi", "Ω": "Omega",
}
FUNCTIONS = {"sin", "cos", "tan", "cot", "sec", "csc", "arcsin", "arccos", "arctan", "sinh",
             "cosh", "tanh", "exp", "log", "ln", "lg", "lim", "max", "min", "sup", "inf", "det",
             "arg", "deg", "dim", "gcd", "ker", "Pr"}
ACCENTS = {"̂": r"\hat", "̃": r"\tilde", "̄": r"\bar", "̅": r"\bar",
           "̇": r"\dot", "̈": r"\ddot", "⃛": r"\dddot", "⃗": r"\vec",
           "̌": r"\check", "̆": r"\breve", "́": r"\acute", "̀": r"\grave"}
NARY = {"∫": r"\int", "∬": r"\iint", "∭": r"\iiint", "∮": r"\oint",
        "∑": r"\sum", "∏": r"\prod", "∐": r"\coprod", "⋃": r"\bigcup",
        "⋂": r"\bigcap"}
DELIMS = {"": ".", "{": r"\{", "}": r"\}", "‖": r"\|", "⟨": r"\langle",
          "⟩": r"\rangle", "⌊": r"\lfloor", "⌋": r"\rfloor", "⌈": r"\lceil",
          "⌉": r"\rceil"}


def _math_char(ch):
    """One Unicode character -> LaTeX math (handles 'Mathematical Italic' letters etc.)."""
    if ch in MATH_SYMBOLS:
        return MATH_SYMBOLS[ch]
    if ch in GREEK:
        return "\\" + GREEK[ch] + " "
    if ord(ch) < 128:
        return ch
    base = unicodedata.normalize("NFKC", ch)
    if base == ch or len(base) != 1:
        return ch
    name = unicodedata.name(ch, "")
    inner = _math_char(base)
    if "SCRIPT" in name or "ℒ" == ch:
        return r"\mathcal{%s}" % inner.strip()
    if "DOUBLE-STRUCK" in name:
        return r"\mathbb{%s}" % inner.strip()
    if "BOLD" in name:
        return r"\boldsymbol{%s}" % inner.strip() if inner.startswith("\\") else r"\mathbf{%s}" % inner
    return inner


def _mval(el, tag, default=None):
    """Value of a property like <m:chr m:val="x"/> inside el (e.g. inside m:dPr)."""
    if el is None:
        return default
    found = el.find(".//" + M + tag)
    if found is None:
        return default
    return found.get(M + "val", default if default is not None else "")


def _m(el):
    """Convert the children of an OMML element to LaTeX."""
    if el is None:
        return ""
    return "".join(_m_node(child) for child in el)


def _arg(el):
    s = _m(el).strip()
    return s if len(s) == 1 and s.isalnum() else "{" + s + "}"


def _m_node(el):
    tag = etree.QName(el).localname
    if tag.endswith("Pr") or tag == "ctrlPr":
        return ""
    if tag == "r":
        text = "".join(t.text or "" for t in el.iter(M + "t"))
        rpr = el.find(M + "rPr")
        plain = rpr is not None and (rpr.find(M + "nor") is not None or _mval(rpr, "sty") == "p")
        if text in FUNCTIONS:
            return "\\" + text + " "
        if plain and len(text) > 1 and text.isalpha():
            return r"\mathrm{%s}" % esc(text)
        if any(c.isalpha() and 0x80 <= ord(c) <= 0x24F for c in text):  # é, à, ... in math
            return r"\text{%s}" % esc(text)
        return "".join(_math_char(c) for c in text)
    if tag == "f":
        num, den = _m(el.find(M + "num")), _m(el.find(M + "den"))
        if _mval(el.find(M + "fPr"), "type") == "lin":
            return "%s/%s" % (_arg(el.find(M + "num")), _arg(el.find(M + "den")))
        return r"\frac{%s}{%s}" % (num.strip(), den.strip())
    if tag == "sSup":
        sup = _m(el.find(M + "sup")).strip()
        if sup and set(sup) == {"'"}:  # f′ is written f' in LaTeX, not f^'
            return _arg(el.find(M + "e")) + sup
        return "%s^%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sup")))
    if tag == "sSub":
        return "%s_%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sub")))
    if tag == "sSubSup":
        return "%s_%s^%s" % (_arg(el.find(M + "e")), _arg(el.find(M + "sub")), _arg(el.find(M + "sup")))
    if tag == "sPre":
        return "{}_%s^%s%s" % (_arg(el.find(M + "sub")), _arg(el.find(M + "sup")), _arg(el.find(M + "e")))
    if tag == "d":
        pr = el.find(M + "dPr")
        beg, end = _mval(pr, "begChr", "("), _mval(pr, "endChr", ")")
        sep = _mval(pr, "sepChr", "|")
        parts = [_m(e).strip() for e in el.findall(M + "e")]
        inner = (" %s " % MATH_SYMBOLS.get(sep, sep).strip()).join(parts)
        return r"\left%s %s \right%s " % (DELIMS.get(beg, beg), inner, DELIMS.get(end, end))
    if tag == "rad":
        deg = _m(el.find(M + "deg")).strip()
        body = _m(el.find(M + "e")).strip()
        return r"\sqrt[%s]{%s}" % (deg, body) if deg else r"\sqrt{%s}" % body
    if tag == "acc":
        cmd = ACCENTS.get(_mval(el.find(M + "accPr"), "chr", "̂"), r"\hat")
        return "%s{%s}" % (cmd, _m(el.find(M + "e")).strip())
    if tag == "bar":
        cmd = r"\underline" if _mval(el.find(M + "barPr"), "pos") == "bot" else r"\overline"
        return "%s{%s}" % (cmd, _m(el.find(M + "e")).strip())
    if tag == "groupChr":
        pr = el.find(M + "groupChrPr")
        chr_ = _mval(pr, "chr", "⏟")
        body = _m(el.find(M + "e")).strip()
        if chr_ == "⏟":
            return r"\underbrace{%s}" % body
        if chr_ == "⏞":
            return r"\overbrace{%s}" % body
        cmd = r"\overset" if _mval(pr, "pos") == "top" else r"\underset"
        return r"%s{%s}{%s}" % (cmd, _math_char(chr_).strip(), body)
    if tag in ("limLow", "limUpp"):
        base = _m(el.find(M + "e")).strip()
        lim = _m(el.find(M + "lim")).strip()
        if base.startswith("\\") and base[1:].strip() in FUNCTIONS:
            return "%s%s{%s} " % (base.strip(), "_" if tag == "limLow" else "^", lim)
        cmd = r"\underset" if tag == "limLow" else r"\overset"
        return r"%s{%s}{%s}" % (cmd, lim, base)
    if tag == "func":
        name = _m(el.find(M + "fName")).strip()
        if name.startswith(r"\mathrm{") and name[8:-1] in FUNCTIONS:
            name = "\\" + name[8:-1]
        return "%s %s" % (name, _m(el.find(M + "e")).strip())
    if tag == "nary":
        pr = el.find(M + "naryPr")
        op = NARY.get(_mval(pr, "chr", "∫"), r"\int")
        s = op
        if _mval(pr, "subHide") not in ("1", "on", "true"):
            sub = _m(el.find(M + "sub")).strip()
            s += "_{%s}" % sub if sub else ""
        if _mval(pr, "supHide") not in ("1", "on", "true"):
            sup = _m(el.find(M + "sup")).strip()
            s += "^{%s}" % sup if sup else ""
        return s + " " + _m(el.find(M + "e")).strip()
    if tag == "eqArr":
        rows = [_m(e).strip() for e in el.findall(M + "e")]
        if not any("&" in r for r in rows):
            rows = [r.replace("=", "&=", 1) for r in rows]
        return "\\begin{aligned}" + r" \\ ".join(rows) + r"\end{aligned}"
    if tag == "m":
        rows = [" & ".join(_m(e).strip() for e in mr.findall(M + "e"))
                for mr in el.findall(M + "mr")]
        return "\\begin{matrix}" + r" \\ ".join(rows) + r"\end{matrix}"
    # m:e, m:box, m:borderBox, m:phant, m:num, ... : just their content
    return _m(el)


def _tidy_math(s):
    s = " ".join(s.split())
    s = s.replace("{ ", "{").replace(" }", "}")
    # Entr\text{é}e -> \text{Entrée}
    return re.sub(r"([A-Za-z]*)\\text\{([^{}\\]*)\}([A-Za-z]*)", r"\\text{\1\2\3}", s)


def omml_to_latex(a14m):
    """Convert an a14:m element to a list of LaTeX strings ($...$ or \\[...\\])."""
    out = []
    for el in a14m:
        tag = etree.QName(el).localname
        if tag == "oMathPara":
            for om in el.iter(M + "oMath"):
                out.append(r"\[ " + _tidy_math(_m(om)) + r" \]")
        elif tag == "oMath":
            out.append("$" + _tidy_math(_m(el)) + "$")
    return out


# --------------------------------------------------------------------------
# Slide analysis helpers
# --------------------------------------------------------------------------

SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
BULLET_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, None}


def _ph_type(shape):
    if not shape.is_placeholder:
        return "none"
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def _is_title_slide(slide):
    return any(_ph_type(s) == PP_PLACEHOLDER.CENTER_TITLE for s in slide.placeholders)


def _bullet_from_ppr(ppr):
    if ppr is None:
        return None
    if ppr.find(qn("a:buNone")) is not None:
        return "none"
    if ppr.find(qn("a:buAutoNum")) is not None:
        return "enumerate"
    if ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buBlip")) is not None:
        return "itemize"
    return None


def _inherited_bullet(shape, level):
    """Look up the bullet style in the shape's list style and its layout/master placeholders."""
    chain = [shape]
    base = getattr(shape, "_base_placeholder", None)
    while base is not None:
        chain.append(base)
        base = getattr(base, "_base_placeholder", None)
    for sh in chain:
        lst = sh._element.find(".//" + qn("a:lstStyle"))
        if lst is not None:
            kind = _bullet_from_ppr(lst.find(qn(f"a:lvl{level + 1}pPr")))
            if kind:
                return kind
    return None


def paragraph_kind(shape, paragraph):
    kind = _bullet_from_ppr(paragraph._p.pPr)
    if kind is None and shape.is_placeholder:
        kind = _inherited_bullet(shape, paragraph.level)
        if kind is None and _ph_type(shape) in BULLET_PLACEHOLDERS:
            kind = "itemize"
    return None if kind in (None, "none") else kind


def _paragraph_children(p):
    """Yield the run-level children of a paragraph, unwrapping mc:AlternateContent."""
    for el in p:
        if etree.QName(el).localname == "AlternateContent":
            choice = el.find(MC + "Choice")
            if choice is not None:
                yield from _paragraph_children(choice)
        else:
            yield el


def paragraph_latex(paragraph, part, line_break=r"\\ "):
    """Paragraph text with bold/italic/hyperlinks/equations, adjacent equal runs merged."""
    pieces = []  # [text, bold, italic, url] for text, or a raw LaTeX string
    for el in _paragraph_children(paragraph._p):
        tag = etree.QName(el).localname
        if tag == "br":
            pieces.append(line_break)
            continue
        if tag == "m":  # a14:m wraps an Office equation
            pieces.extend(omml_to_latex(el))
            continue
        if tag not in ("r", "fld"):
            continue
        text = "".join(t.text or "" for t in el.iter(qn("a:t")))
        rpr = el.find(qn("a:rPr"))
        bold = rpr is not None and rpr.get("b") in ("1", "true")
        italic = rpr is not None and rpr.get("i") in ("1", "true")
        url = None
        link = rpr.find(qn("a:hlinkClick")) if rpr is not None else None
        if link is not None and link.get(qn("r:id")):
            try:
                url = part.rels[link.get(qn("r:id"))].target_ref
            except KeyError:
                url = None
        if pieces and isinstance(pieces[-1], list) and pieces[-1][1:] == [bold, italic, url]:
            pieces[-1][0] += text
        else:
            pieces.append([text, bold, italic, url])

    out = []
    for piece in pieces:
        if isinstance(piece, str):
            out.append(piece)
            continue
        text, bold, italic, url = piece
        s = esc(text)
        if s.strip():
            if bold:
                s = r"\textbf{" + s + "}"
            if italic:
                s = r"\textit{" + s + "}"
            if url:
                s = r"\href{" + esc_url(url) + "}{" + s + "}"
        out.append(s)
    return "".join(out).strip()


def render_text_frame(shape, part):
    """Turn a text frame into itemize/enumerate blocks and plain paragraphs."""
    lines, stack = [], []  # stack of (env, level)

    def close_to(level):
        while stack and stack[-1][1] > level:
            lines.append("  " * (len(stack) - 1) + r"\end{" + stack.pop()[0] + "}")

    for p in shape.text_frame.paragraphs:
        text = paragraph_latex(p, part)
        if not text:
            continue
        kind = paragraph_kind(shape, p)
        if kind is None:
            close_to(-1)
            if lines:
                lines.append("")
            lines.append(text)
            continue
        level = min(p.level, 3)
        close_to(level)
        if stack and stack[-1][1] == level and stack[-1][0] != kind:
            close_to(level - 1)
        if not stack or stack[-1][1] < level:
            lines.append("  " * len(stack) + r"\begin{" + kind + "}")
            stack.append((kind, level))
        lines.append("  " * len(stack) + r"\item " + text)
    close_to(-1)
    return "\n".join(lines)


def render_table(table, part):
    ncols = len(table.columns)
    rows = [r"\begin{tabular}{|" + "l|" * ncols + "}", r"\hline"]
    for row in table.rows:
        cells, skip = [], 0
        for cell in row.cells:
            if skip:
                skip -= 1
                continue
            text = " ".join(filter(None, (paragraph_latex(p, part, " ")
                                          for p in cell.text_frame.paragraphs)))
            if cell.is_spanned:
                text = ""
            if cell.is_merge_origin and cell.span_width > 1:
                skip = cell.span_width - 1
                text = r"\multicolumn{%d}{|l|}{%s}" % (cell.span_width, text)
            cells.append(text)
        rows.append(" & ".join(cells) + r" \\ \hline")
    rows.append(r"\end{tabular}")
    body = "\n".join(rows)
    if ncols > 5:
        body = "\\resizebox{\\linewidth}{!}{%\n" + body + "}"
    return body


# --------------------------------------------------------------------------
# Slide conversion
# --------------------------------------------------------------------------

class Element:
    """A piece of slide content with its position (EMU, slide coordinates).

    kind: text | table | image | drawing | comment
    ids:  ids of the top-level shapes it comes from (used to render diagrams)
    """

    def __init__(self, kind, x, y, w, h, tex=None, ids=(), movable=False):
        self.kind, self.x, self.y, self.w, self.h = kind, x, y, w, h
        self.tex = tex
        self.ids = set(ids)
        self.movable = movable     # may be absorbed into a diagram (text box / picture)
        self.src = None            # (blob, content_type, crop) for pictures
        self.render = None         # render job index for diagrams
        self.strong = False        # group or connector: always a diagram
        self.fallback = []         # what to output if the diagram cannot be rendered
        self.image = None          # (name, ok, filename) once saved

    @property
    def right(self):
        return self.x + self.w

    @property
    def bottom(self):
        return self.y + self.h


def image_order(elements):
    """Image elements in reading order; position k is saved as slide_<i>_image_<k+1>."""
    return [e for e in sorted(elements, key=lambda e: (e.y, e.x)) if e.kind == "image"]


def _union_box(elements):
    x0 = min(e.x for e in elements)
    y0 = min(e.y for e in elements)
    return x0, y0, max(e.right for e in elements) - x0, max(e.bottom for e in elements) - y0


class SlideConverter:
    def __init__(self, prs, images_dir, warnings, diagrams=True):
        self.prs = prs
        self.sw, self.sh = prs.slide_width, prs.slide_height
        self.images_dir = images_dir
        self.warnings = warnings
        self.diagrams = diagrams
        self.jobs = []  # diagrams to render: (slide index, shape ids, bbox)

    # ---- collecting elements ------------------------------------------------

    def collect(self, container, slide_no, part, transform, out, skip_ids, top_id=None):
        for shape in iter_shapes(container, self.shapes):
            if shape.shape_id in skip_ids:
                continue
            sid = top_id if top_id is not None else shape.shape_id
            nested = top_id is not None
            try:
                x, y, w, h = transform(shape.left or 0, shape.top or 0,
                                       shape.width or 0, shape.height or 0)
            except TypeError:
                x = y = w = h = 0
            stype = shape.shape_type

            def add(kind, tex=None, movable=False):
                el = Element(kind, x, y, w, h, tex, ids={sid}, movable=movable)
                out.append(el)
                return el

            if stype == MSO_SHAPE_TYPE.GROUP:
                children = []
                self.collect(shape._element, slide_no, part,
                             self._group_transform(shape, transform), children, skip_ids, sid)
                if nested:
                    out.extend(children)
                else:
                    el = add("drawing")
                    el.strong, el.fallback = True, children
                continue
            if shape.is_placeholder and _ph_type(shape) in SKIP_PLACEHOLDERS:
                continue

            if stype == MSO_SHAPE_TYPE.MEDIA or hasattr(shape, "poster_frame"):
                poster = getattr(shape, "poster_frame", None)
                if shape._element.find(".//" + qn("a:audioFile")) is not None:
                    add("comment", "% skipped: audio")
                elif poster is None:
                    add("comment", "% skipped: video without poster frame")
                else:
                    el = add("image", "% video poster frame")
                    el.src = (poster.blob, poster.content_type, (0, 0, 0, 0))
                continue

            if _has_image(shape):
                if shape._element.find(".//" + qn("a:audioFile")) is not None:
                    add("comment", "% skipped: audio")
                    continue
                crop = tuple(max(0.0, getattr(shape, a, 0) or 0) for a in
                             ("crop_left", "crop_top", "crop_right", "crop_bottom"))
                el = add("image", movable=not shape.is_placeholder)
                el.src = (shape.image.blob, shape.image.content_type, crop)
                continue

            if getattr(shape, "has_table", False) and shape.has_table:
                add("table", render_table(shape.table, part))
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                add("comment", "% skipped: chart")
                self.warnings.append(f"slide {slide_no}: chart skipped")
                continue

            is_line = (stype == MSO_SHAPE_TYPE.LINE
                       or etree.QName(shape._element).localname == "cxnSp")
            is_autoshape = (stype in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM)
                            and not shape.is_placeholder)
            if not nested and (is_line or is_autoshape):
                el = add("drawing")
                el.strong = is_line
                tex = render_text_frame(shape, part) if shape.has_text_frame else ""
                if tex:
                    el.fallback = [Element("text", x, y, w, h, tex, ids={sid})]
                continue
            if is_line:
                continue

            if shape.has_text_frame:
                tex = render_text_frame(shape, part)
                if tex:
                    add("text", tex, movable=not shape.is_placeholder)
                continue
            if stype in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                add("comment", "% skipped: embedded object")
            elif etree.QName(shape._element).localname == "graphicFrame":
                add("comment", "% skipped: SmartArt/diagram")
                self.warnings.append(f"slide {slide_no}: SmartArt/diagram skipped")

    @staticmethod
    def _group_transform(group, parent):
        xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
        ox, oy = int(off.get("x")), int(off.get("y"))
        ex, ey = int(ext.get("cx")), int(ext.get("cy"))
        cx, cy = int(choff.get("x")), int(choff.get("y"))
        cex, cey = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
        sx, sy = ex / cex, ey / cey

        def transform(x, y, w, h):
            return parent(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy)
        return transform

    # ---- diagrams -----------------------------------------------------------

    def group_diagrams(self, elements, slide_index):
        """Merge touching drawing shapes (and the labels on top of them) into diagrams."""
        margin = 0.03 * min(self.sw, self.sh)
        drawings = [e for e in elements if e.kind == "drawing"]
        rest = [e for e in elements if e.kind != "drawing"]

        def near(a, b, m):
            return (a.x - m < b.right and b.x - m < a.right and
                    a.y - m < b.bottom and b.y - m < a.bottom)

        clusters = []
        for d in drawings:
            hits = [c for c in clusters if any(near(d, o, margin / 3) for o in c)]
            merged = [d] + [o for c in hits for o in c]
            clusters = [c for c in clusters if all(c is not h for h in hits)] + [merged]

        result = []
        for c in clusters:
            if not self.diagrams or not (len(c) >= 2 or c[0].strong):
                result.extend(f for d in c for f in d.fallback)  # keep the text only
                continue
            x, y, w, h = _union_box(c)
            box = Element("drawing", x - margin, y - margin, w + 2 * margin, h + 2 * margin)
            absorbed = [e for e in rest if e.movable and
                        box.x < e.x + e.w / 2 < box.right and box.y < e.y + e.h / 2 < box.bottom]
            rest = [e for e in rest if all(e is not a for a in absorbed)]
            members = c + absorbed
            el = Element("image", *_union_box(members), ids=set().union(*(m.ids for m in members)))
            el.fallback = [f for d in c for f in d.fallback] + absorbed
            el.render = len(self.jobs)
            self.jobs.append((slide_index, el.ids, (el.x, el.y, el.w, el.h)))
            result.append(el)
        return rest + result

    # ---- output ---------------------------------------------------------------

    def name_images(self, elements, slide_no, rendered):
        """Save pictures / rendered diagrams as slide_<i>_image_<j> in reading order."""
        final = []
        for el in elements:
            if el.render is not None and el.render not in rendered:
                final.extend(el.fallback)  # the diagram could not be rendered
            else:
                final.append(el)
        for count, el in enumerate(image_order(final), start=1):
            name = f"slide_{slide_no}_image_{count}"
            if el.render is not None:
                fname = name + ".png"
                shutil.move(rendered[el.render], os.path.join(self.images_dir, fname))
                el.image = (name, True, fname)
                continue
            fname, ok = save_image(el.src[0], el.src[1], self.images_dir, name, el.src[2])
            if not ok:
                self.warnings.append(f"slide {slide_no}: could not convert {fname} to PNG "
                                     "(install Inkscape or LibreOffice, or convert it by hand)")
            el.image = (name, ok, fname)
        return final

    def image_tex(self, el, box_w):
        name, ok, fname = el.image
        wf = min(el.w / box_w, 1.0) if box_w else 1.0
        hf = el.h / self.sh if self.sh else 0.5
        if hf > 0.6 or (hf > 0.45 and el.h > el.w):
            size = "height=%.2f\\textheight" % min(hf, 0.75)
        else:
            size = "width=%.2f\\linewidth" % max(wf, 0.1)
        line = r"\includegraphics[%s]{%s}" % (size, name)
        if not ok:
            line = f"% TODO convert {fname} to PNG, then uncomment:\n% " + line
        return line

    def block_tex(self, elements, box_w):
        parts = []
        for el in sorted(elements, key=lambda e: (e.y, e.x)):
            if el.kind == "image":
                img = self.image_tex(el, box_w)
                if el.tex:
                    img = el.tex + "\n" + img
                parts.append("\\begin{center}\n" + img + "\n\\end{center}")
            elif el.kind == "table":
                parts.append("\\begin{center}\n" + el.tex + "\n\\end{center}")
            else:
                parts.append(el.tex)
        return "\n\n".join(parts)

    @staticmethod
    def split_columns(elements):
        """Group elements into side-by-side columns (by horizontal overlap)."""
        cols = []
        tol = 0.02
        for el in sorted(elements, key=lambda e: e.x):
            if cols and el.x < max(e.right for e in cols[-1]) - tol * max(el.w, 1):
                cols[-1].append(el)
            else:
                cols.append([el])
        return cols

    def body_tex(self, elements):
        content = [e for e in elements if e.kind != "comment"]
        comments = [e.tex for e in elements if e.kind == "comment"]

        # a picture covering (almost) the whole slide behind other content is a background
        full = [e for e in content if e.kind == "image" and e.w * e.h > 0.85 * self.sw * self.sh]
        if full and len(content) > len(full):
            for e in full:
                comments.append(f"% background image not included: {e.image[2]}")
                content.remove(e)

        cols = self.split_columns(content)
        if 2 <= len(cols) <= 4:
            spans = [max(e.right for e in c) - min(e.x for e in c) for c in cols]
            total = sum(spans) or 1
            out = ["\\begin{columns}"]
            for col, span in zip(cols, spans):
                frac = max(0.15, round(0.95 * span / total, 2))
                out.append("\\begin{column}{%.2f\\textwidth}" % frac)
                out.append(self.block_tex(col, span))
                out.append("\\end{column}")
            out.append("\\end{columns}")
            body = "\n".join(out)
        else:
            body = self.block_tex(content, self.sw)
        return "\n".join(comments + ([body] if body else []))

    def prepare(self, slide, slide_index, is_first_title):
        """First pass: collect the slide's content; diagrams are queued for rendering."""
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = " ".join(title_shape.text_frame.text.replace("\v", " ").split())

        skip = {title_shape.shape_id} if title_shape is not None else set()
        subtitle = ""
        if is_first_title:
            for ph in slide.placeholders:
                if _ph_type(ph) == PP_PLACEHOLDER.SUBTITLE:
                    subtitle = " ".join(ph.text_frame.text.replace("\v", " ").split())
                    skip.add(ph.shape_id)

        elements = []
        identity = lambda x, y, w, h: (x, y, w, h)
        self.shapes = slide.shapes
        self.collect(slide.shapes._spTree, slide_index + 1, slide.part, identity, elements, skip)
        if is_first_title:  # logos etc. on the title slide are only extracted
            elements = [e for e in elements if e.kind == "image"]
        else:
            elements = self.group_diagrams(elements, slide_index)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        return dict(title=title, subtitle=subtitle, elements=elements, notes=notes,
                    first_title=is_first_title)

    def finish(self, state, slide_no, rendered):
        """Second pass: save the images and write the frame."""
        elements = self.name_images(state["elements"], slide_no, rendered)
        lines = []
        if state["first_title"]:
            lines.append("\\begin{frame}")
            lines.append("\\titlepage")
            for e in elements:
                lines.append(f"% not included from title slide: {e.image[2]}")
        else:
            title = state["title"]
            lines.append("\\begin{frame}{%s}" % esc(title) if title else "\\begin{frame}")
            body = self.body_tex(elements)
            if body:
                lines.append(body)
        if state["notes"]:
            lines.append("\\note{%s}" % esc(state["notes"]).replace("\n", "\n\n"))
        lines.append("\\end{frame}")
        return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------
# Rendering diagrams with PowerPoint (Windows)
# --------------------------------------------------------------------------

RENDER_PS1 = r"""
param([string]$pptx, [string]$outdir, [string]$slides, [int]$width, [int]$height)
$ErrorActionPreference = "Stop"
$app = New-Object -ComObject PowerPoint.Application
$pres = $null
try {
    $pres = $app.Presentations.Open($pptx, -1, 0, 0)
    foreach ($n in $slides.Split(",")) {
        $pres.Slides.Item([int]$n).Export((Join-Path $outdir "s$n.png"), "PNG", $width, $height)
    }
} finally {
    if ($pres) { $pres.Close() }
    if ($app.Presentations.Count -eq 0) { $app.Quit() }
}
"""


def _element_id(el):
    cnvpr = el.find(".//" + qn("p:cNvPr"))
    return int(cnvpr.get("id")) if cnvpr is not None else None


def render_diagrams(pptx_path, jobs, background, tmpdir, warnings):
    """Render each job (slide index, shape ids, bbox) to a cropped PNG using PowerPoint.

    Works on a temporary copy of the deck in which every slide keeps only its diagram
    shapes, so nothing else ends up in the picture. Returns {job index: png path}.
    """
    if not jobs:
        return {}
    if os.name != "nt" or not shutil.which("powershell") or Image is None:
        warnings.append("diagrams kept as text: rendering them needs PowerPoint on Windows")
        return {}

    prs = Presentation(pptx_path)
    keep = {}
    for slide_index, ids, _ in jobs:
        keep.setdefault(slide_index, set()).update(ids)
    bg_xml = ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr>'
              '<a:solidFill><a:srgbClr val="%02X%02X%02X"/></a:solidFill><a:effectLst/>'
              '</p:bgPr></p:bg>' % tuple(background))
    for index, slide in enumerate(prs.slides):
        sld = slide._element
        sld.attrib.pop("show", None)
        sld.set("showMasterSp", "0")
        tree = slide.shapes._spTree
        for child in list(tree):
            if etree.QName(child).localname in ("nvGrpSpPr", "grpSpPr", "extLst"):
                continue
            if _element_id(child) not in keep.get(index, ()):
                tree.remove(child)
        csld = sld.find(qn("p:cSld"))
        old = csld.find(qn("p:bg"))
        if old is not None:
            csld.remove(old)
        csld.insert(0, etree.fromstring(bg_xml))
    deck = os.path.join(tmpdir, "render.pptx")
    prs.save(deck)

    width = 2400
    height = round(width * prs.slide_height / prs.slide_width)
    script = os.path.join(tmpdir, "render.ps1")
    with open(script, "w", encoding="utf-8") as f:
        f.write(RENDER_PS1)
    slides = ",".join(str(i + 1) for i in sorted(keep))
    print(f"Rendering {len(jobs)} diagram(s) with PowerPoint...")
    try:
        proc = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                               "-File", script, "-pptx", os.path.abspath(deck),
                               "-outdir", os.path.abspath(tmpdir), "-slides", slides,
                               "-width", str(width), "-height", str(height)],
                              capture_output=True, text=True, timeout=600)
    except (OSError, subprocess.SubprocessError) as e:
        warnings.append(f"diagrams kept as text: PowerPoint rendering failed ({e})")
        return {}
    if proc.returncode != 0:
        msg = (proc.stderr or proc.stdout).strip().splitlines()
        warnings.append("diagrams kept as text: PowerPoint rendering failed (%s)"
                        % (msg[0] if msg else "unknown error"))
        return {}

    out = {}
    scale = width / prs.slide_width
    pad = int(0.05 * width)
    per_slide = {}
    for slide_index, _, _ in jobs:
        per_slide[slide_index] = per_slide.get(slide_index, 0) + 1
    for key, (slide_index, _, (x, y, w, h)) in enumerate(jobs):
        png = os.path.join(tmpdir, f"s{slide_index + 1}.png")
        if not os.path.exists(png):
            continue
        with Image.open(png) as im:
            im = im.convert("RGB")
            if per_slide[slide_index] == 1:
                # only this diagram is on the rendered slide: trim below finds its extent
                # (connectors may be drawn outside their bounding box)
                crop = im.copy()
            else:
                crop = im.crop((max(0, int(x * scale) - pad), max(0, int(y * scale) - pad),
                                min(im.width, int((x + w) * scale) + pad),
                                min(im.height, int((y + h) * scale) + pad)))
        # trim to the drawn content, keeping a small border
        diff = ImageChops.difference(crop, Image.new("RGB", crop.size, tuple(background)))
        content = diff.convert("L").point(lambda v: 255 if v > 12 else 0).getbbox()
        if content:
            b = int(0.01 * width)
            crop = crop.crop((max(0, content[0] - b), max(0, content[1] - b),
                              min(crop.width, content[2] + b), min(crop.height, content[3] + b)))
        path = os.path.join(tmpdir, f"job{key}.png")
        crop.save(path)
        out[key] = path
    return out


SHAPE_TAGS = {"sp", "pic", "grpSp", "graphicFrame", "cxnSp"}


def iter_shapes(container, parent):
    """Shapes in a shape tree or group, including those wrapped in mc:AlternateContent
    (PowerPoint uses it for text boxes containing equations; python-pptx skips them)."""
    for child in container:
        tag = etree.QName(child).localname
        if tag == "AlternateContent":
            choice = child.find(MC + "Choice")
            if choice is not None:
                yield from iter_shapes(choice, parent)
        elif tag in SHAPE_TAGS:
            yield SlideShapeFactory(child, parent)


def _has_image(shape):
    if shape.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return False
    try:
        return shape.image is not None
    except (AttributeError, ValueError, KeyError):
        return False


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def build_main(prs, theme, color_line, n_slides, hidden, title, subtitle):
    ratio = prs.slide_width / prs.slide_height
    aspect = "169" if abs(ratio - 16 / 9) < 0.05 else "1610" if abs(ratio - 1.6) < 0.03 else "43"
    author = prs.core_properties.author or ""
    lines = [
        "\\documentclass[aspectratio=%s]{beamer}" % aspect,
        "\\usetheme{%s}" % theme,
    ]
    lines += color_line
    lines += [
        "\\graphicspath{{images/}}",
        "",
        "\\title{%s}" % esc(title),
    ]
    if subtitle:
        lines.append("\\subtitle{%s}" % esc(subtitle))
    lines += ["\\author{%s}" % esc(author), "", "\\begin{document}", ""]
    for i in range(1, n_slides + 1):
        prefix = "% (hidden) " if i in hidden else ""
        lines.append(prefix + "\\input{slides/slide_%d}" % i)
    lines += ["", "\\end{document}", ""]
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convert a PowerPoint (.pptx) file to a Beamer project.")
    ap.add_argument("pptx", nargs="?", help="input .pptx file")
    ap.add_argument("-o", "--output", help="output folder (default: <name>_beamer)")
    ap.add_argument("--theme", default="Madrid", help="Beamer theme, e.g. Madrid, Warsaw, metropolis")
    ap.add_argument("--color", default="auto",
                    help="colour theme: auto (closest to the deck), exact (deck colour), "
                         "none, or a Beamer colour theme name such as beaver")
    ap.add_argument("--no-render", action="store_true",
                    help="do not render drawn diagrams with PowerPoint; keep their text only")
    ap.add_argument("--list-themes", action="store_true", help="list known themes and exit")
    args = ap.parse_args(argv)

    if args.list_themes:
        print("Themes:       " + ", ".join(THEMES))
        print("Colour themes: " + ", ".join(list(COLOR_THEMES) + OTHER_COLOR_THEMES))
        return 0
    if not args.pptx:
        ap.error("the input .pptx file is required")

    prs = Presentation(args.pptx)
    out_dir = args.output or os.path.splitext(args.pptx)[0] + "_beamer"
    slides_dir = os.path.join(out_dir, "slides")
    images_dir = os.path.join(out_dir, "images")
    os.makedirs(slides_dir, exist_ok=True)
    os.makedirs(images_dir, exist_ok=True)

    # colours
    accent, background = deck_colors(prs)
    exact = "\\usecolortheme[RGB={%d,%d,%d}]{structure}" % accent
    if args.color == "exact":
        color_line = [exact]
    elif args.color == "none":
        color_line = []
    else:
        name = pick_color_theme(accent, background) if args.color == "auto" else args.color
        color_line = ["\\usecolortheme{%s}" % name,
                      "%% exact deck colour instead: %s" % exact]

    warnings = []
    conv = SlideConverter(prs, images_dir, warnings, diagrams=not args.no_render)
    title = prs.core_properties.title or os.path.splitext(os.path.basename(args.pptx))[0]
    subtitle = ""
    hidden = set()

    # pass 1: read every slide, queue diagrams for rendering
    states = []
    for index, slide in enumerate(prs.slides):
        if slide._element.get("show") in ("0", "false"):
            hidden.add(index + 1)
        first_title = not any(s["first_title"] for s in states) and _is_title_slide(slide)
        states.append(conv.prepare(slide, index, first_title))
        if first_title:
            title = states[-1]["title"] or title
            subtitle = states[-1]["subtitle"]

    # render diagrams, then pass 2: write the slides
    with tempfile.TemporaryDirectory() as tmp:
        rendered = render_diagrams(args.pptx, conv.jobs, background, tmp, warnings)
        for n, state in enumerate(states, start=1):
            with open(os.path.join(slides_dir, f"slide_{n}.tex"), "w", encoding="utf-8") as f:
                f.write(conv.finish(state, n, rendered))
    n = len(states)

    with open(os.path.join(out_dir, "main.tex"), "w", encoding="utf-8") as f:
        f.write(build_main(prs, args.theme, color_line, n, hidden, title, subtitle))

    print(f"Converted {n} slides -> {out_dir}")
    print("Deck accent colour: RGB%s, background: RGB%s" % (accent, background))
    if color_line:
        print("Colour: " + color_line[0])
    for w in warnings:
        print("warning: " + w, file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())

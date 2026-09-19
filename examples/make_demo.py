#!/usr/bin/env python3
"""Build examples/demo.pptx, a small deck that shows every feature of PowerPoint2Beamer.

    python examples/make_demo.py
    python powerpoint2beamer.py examples/demo.pptx -o examples/demo_beamer --compile
    python powerpoint2tikz.py examples/demo.pptx --slides 6 --project examples/demo_beamer

Author: Abolfazl Mohebbi, PhD., P. Eng. (Polytechnique Montréal) - MIT License
"""

import io
import math
import os
import uuid

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
RED, GREEN = RGBColor(0xC0, 0x00, 0x00), RGBColor(0x00, 0xB0, 0x50)


# ---------------------------------------------------------------- helpers

def math_run(text):
    return ('<m:r><a:rPr lang="en-US" sz="2400" i="1"><a:latin typeface="Cambria Math"/></a:rPr>'
            '<m:t>%s</m:t></m:r>' % text)


def equation_box(slide, sid, x, y, w, h, paragraphs):
    """A text box with Office equations. paragraphs: ('text', str) or ('math', omml)."""
    body = []
    for kind, content in paragraphs:
        if kind == "text":
            body.append('<a:p><a:r><a:rPr lang="en-US" sz="2000"/><a:t>%s</a:t></a:r></a:p>' % content)
        else:
            body.append('<a:p><a:pPr algn="ctr"/><a14:m><m:oMathPara><m:oMath>%s</m:oMath>'
                        '</m:oMathPara></a14:m></a:p>' % content)
    sp = ('<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="Equations"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
          '</p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
          '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr><p:txBody>'
          '<a:bodyPr wrap="square"><a:spAutoFit/></a:bodyPr><a:lstStyle/>{body}</p:txBody></p:sp>')
    choice = sp.format(sid=sid, x=x, y=y, w=w, h=h, body="".join(body))
    fallback = sp.format(sid=sid, x=x, y=y, w=w, h=h,
                         body='<a:p><a:r><a:rPr lang="en-US"/><a:t>(equations)</a:t></a:r></a:p>')
    xml = ('<mc:AlternateContent xmlns:mc="{mc}" xmlns:p="{p}" xmlns:a="{a}" xmlns:m="{m}" '
           'xmlns:a14="{a14}"><mc:Choice Requires="a14">{c}</mc:Choice>'
           '<mc:Fallback>{f}</mc:Fallback></mc:AlternateContent>').format(c=choice, f=fallback, **NS)
    slide.shapes._spTree.append(etree.fromstring(xml))


def add_sections(prs, groups):
    ids = [s.get("id") for s in prs.slides._sldIdLst]
    sections = "".join(
        '<p14:section name="%s" id="{%s}"><p14:sldIdLst>%s</p14:sldIdLst></p14:section>'
        % (name, str(uuid.uuid4()).upper(), "".join('<p14:sldId id="%s"/>' % ids[i] for i in members))
        for name, members in groups)
    xml = ('<p:extLst xmlns:p="{p}"><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">'
           '<p14:sectionLst xmlns:p14="{p14}">{s}</p14:sectionLst></p:ext></p:extLst>'
           ).format(s=sections, **NS)
    prs.part._element.append(etree.fromstring(xml))


def appear_on_clicks(slide, spid, paragraphs):
    """'Appear' animation: the given paragraphs of shape spid show up one click at a time."""
    steps, n = [], 3
    for para in paragraphs:
        steps.append(
            '<p:par><p:cTn id="{0}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            '<p:childTnLst><p:par><p:cTn id="{1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst><p:par><p:cTn id="{2}" presetID="1" presetClass="entr" presetSubtype="0" '
            'fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst><p:set><p:cBhvr><p:cTn id="{3}" dur="1" fill="hold"><p:stCondLst>'
            '<p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{4}"><p:txEl>'
            '<p:pRg st="{5}" end="{5}"/></p:txEl></p:spTgt></p:tgtEl><p:attrNameLst>'
            '<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to>'
            '<p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst>'
            '</p:cTn></p:par></p:childTnLst></p:cTn></p:par>'.format(n, n + 1, n + 2, n + 3, spid, para))
        n += 4
    xml = ('<p:timing xmlns:p="{p}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
           'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" '
           'dur="indefinite" nodeType="mainSeq"><p:childTnLst>{s}</p:childTnLst></p:cTn><p:prevCondLst>'
           '<p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
           '</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>'
           '<p:bldP spid="{sp}" grpId="0" build="p"/></p:bldLst></p:timing>').format(
        s="".join(steps), sp=spid, **NS)
    slide._element.append(etree.fromstring(xml))


def arrow(connector):
    ln = connector.line._get_or_add_ln()
    ln.append(etree.fromstring('<a:tailEnd xmlns:a="%s" type="triangle"/>' % NS["a"]))
    connector.line.color.rgb = RGBColor(0, 0, 0)
    connector.line.width = Pt(1.5)


def label(slide, text, x, y, w=0.9, italic=True, size=20):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text, run.font.italic, run.font.size = text, italic, Pt(size)
    return tb


def plot_picture():
    """A small step-response plot, drawn with Pillow, used as a 'picture'."""
    im = Image.new("RGB", (640, 480), "white")
    d = ImageDraw.Draw(im)
    d.line([(60, 420), (620, 420)], fill="black", width=3)
    d.line([(60, 420), (60, 30)], fill="black", width=3)
    pts = []
    for i in range(561):  # underdamped second-order step response
        t = i / 80
        y = 1 - math.exp(-0.5 * t) * (math.cos(1.5 * t) + 0.5 / 1.5 * math.sin(1.5 * t))
        pts.append((60 + i, 420 - 250 * y))
    d.line(pts, fill=(31, 119, 180), width=5)
    d.line([(60, 170), (620, 170)], fill=(200, 200, 200), width=2)
    d.text((300, 440), "time", fill="black")
    d.text((10, 60), "y(t)", fill="black")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- the deck

def build(path):
    prs = Presentation()  # 4:3 Office theme
    layout_title, layout_body, layout_only = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    # 1 title
    s = prs.slides.add_slide(layout_title)
    s.shapes.title.text = "Control Systems in a Nutshell"
    s.placeholders[1].text = "A PowerPoint2Beamer demo"
    prs.core_properties.author = "Abolfazl Mohebbi"

    # 2 text formatting
    s = prs.slides.add_slide(layout_body)
    s.shapes.title.text = "Feedback control"
    tf = s.placeholders[1].text_frame
    p = tf.paragraphs[0]
    p.text = "A controller compares the "
    r = p.add_run(); r.text = "output"; r.font.bold = True
    r = p.add_run(); r.text = " with the "
    r = p.add_run(); r.text = "reference"; r.font.italic = True
    for level, text in [(1, "Proportional, integral and derivative actions"),
                        (1, "Tuned from a model or from experiments"),
                        (0, "Too much gain makes the loop ")]:
        p = tf.add_paragraph(); p.text = text; p.level = level
    r = p.add_run(); r.text = "unstable!"; r.font.bold = True; r.font.color.rgb = RED
    p = tf.add_paragraph(); p.text = "A well-tuned loop is "
    r = p.add_run(); r.text = "fast and stable"; r.font.color.rgb = GREEN

    # 3 equations
    s = prs.slides.add_slide(layout_only)
    s.shapes.title.text = "Transfer functions"
    first_order = math_run("G(s)=") + "<m:f><m:num>%s</m:num><m:den>%s</m:den></m:f>" % (
        math_run("K"), math_run("τs+1"))
    laplace = (math_run("F(s)=") +
               '<m:nary><m:naryPr><m:chr m:val="∫"/></m:naryPr><m:sub>%s</m:sub><m:sup>%s</m:sup>'
               '<m:e>%s<m:sSup><m:e>%s</m:e><m:sup>%s</m:sup></m:sSup>%s</m:e></m:nary>'
               % (math_run("0"), math_run("∞"), math_run("f(t)"), math_run("e"), math_run("−st"),
                  math_run("dt")))
    wn2 = '<m:sSubSup><m:e>%s</m:e><m:sub>%s</m:sub><m:sup>%s</m:sup></m:sSubSup>' % (
        math_run("ω"), math_run("n"), math_run("2"))
    s2 = '<m:sSup><m:e>%s</m:e><m:sup>%s</m:sup></m:sSup>' % (math_run("s"), math_run("2"))
    wn = '<m:sSub><m:e>%s</m:e><m:sub>%s</m:sub></m:sSub>' % (math_run("ω"), math_run("n"))
    second = math_run("H(s)=") + "<m:f><m:num>%s</m:num><m:den>%s</m:den></m:f>" % (
        wn2, s2 + math_run("+2ζ") + wn + math_run("s+") + wn2)
    equation_box(s, 20, Inches(0.5), Inches(1.5), Inches(9), Inches(4.8), [
        ("text", "The Laplace transform turns differential equations into algebra:"),
        ("math", laplace),
        ("text", "First-order and second-order systems:"),
        ("math", first_order),
        ("math", second)])

    # 4 text + picture
    s = prs.slides.add_slide(layout_body)
    s.shapes.title.text = "Step response"
    tf = s.placeholders[1].text_frame
    tf.text = "The output after a sudden change of the input"
    for t in ("Rise time", "Overshoot", "Settling time", "Steady-state error"):
        p = tf.add_paragraph(); p.text = t; p.level = 1
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.5), Inches(1.6), Inches(9), Inches(4.5)
    s.shapes.add_picture(plot_picture(), Inches(5.3), Inches(2.2), width=Inches(4.2))

    # 5 table
    s = prs.slides.add_slide(layout_only)
    s.shapes.title.text = "Ziegler–Nichols tuning"
    rows = [("Controller", "Kp", "Ti", "Td"), ("P", "0.5 Ku", "—", "—"),
            ("PI", "0.45 Ku", "Tu / 1.2", "—"), ("PID", "0.6 Ku", "Tu / 2", "Tu / 8")]
    t = s.shapes.add_table(4, 4, Inches(1), Inches(2), Inches(8), Inches(2.4)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v

    # 6 block diagram (drawn with shapes and connectors)
    s = prs.slides.add_slide(layout_only)
    s.shapes.title.text = "Closed-loop block diagram"
    y = Inches(3.2)
    summ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.0), y - Inches(0.25), Inches(0.5), Inches(0.5))
    ctrl = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.3), y - Inches(0.45), Inches(1.4), Inches(0.9))
    plant = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), y - Inches(0.45), Inches(1.4), Inches(0.9))
    for shape, text in ((ctrl, "C(s)"), (plant, "G(s)")):
        shape.fill.background()
        shape.line.color.rgb = RGBColor(0, 0, 0)
        run = shape.text_frame.paragraphs[0].add_run()
        run.text, run.font.italic, run.font.size = text, True, Pt(20)
        run.font.color.rgb = RGBColor(0, 0, 0)
    summ.fill.background()
    summ.line.color.rgb = RGBColor(0, 0, 0)
    r_label = label(s, "R(s)", 0.6, 2.95)
    y_label = label(s, "Y(s)", 8.5, 2.95)
    label(s, "+", 1.75, 2.45, 0.4, italic=False, size=16)
    label(s, "−", 1.75, 3.35, 0.4, italic=False, size=16)
    for a, b, ai, bi in ((summ, ctrl, 3, 1), (ctrl, plant, 3, 1)):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        c.begin_connect(a, ai)
        c.end_connect(b, bi)
        arrow(c)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(r_label, 3)
    c.end_connect(summ, 1)
    arrow(c)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(plant, 3)
    c.end_connect(y_label, 1)
    arrow(c)
    # feedback: out of G(s), down, back left and up into the summing junction
    fb = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
    fb.begin_connect(plant, 3)
    fb.end_connect(summ, 2)
    geom = fb._element.find(".//" + qn("a:prstGeom"))
    geom.set("prst", "bentConnector4")
    width = fb.width or Emu(1)
    height = fb.height or Emu(1)
    av = geom.find(qn("a:avLst"))
    av.append(etree.fromstring('<a:gd xmlns:a="%s" name="adj1" fmla="val %d"/>'
                               % (NS["a"], -int(100000 * Inches(0.6) / width))))
    av.append(etree.fromstring('<a:gd xmlns:a="%s" name="adj2" fmla="val %d"/>'
                               % (NS["a"], int(100000 * Inches(1.3) / height))))
    arrow(fb)

    # 7 native chart
    s = prs.slides.add_slide(layout_only)
    s.shapes.title.text = "Response to a step input"
    data = XyChartData()

    for name, zeta in (("ζ = 0.3", 0.3), ("ζ = 0.7", 0.7), ("ζ = 1.0", 1.0)):
        series = data.add_series(name)
        for i in range(0, 41):
            t = i * 0.25
            if zeta < 1:
                wd = math.sqrt(1 - zeta ** 2)
                yv = 1 - math.exp(-zeta * t) * (math.cos(wd * t) + zeta / wd * math.sin(wd * t))
            else:
                yv = 1 - math.exp(-t) * (1 + t)
            series.add_data_point(round(t, 2), round(yv, 3))
    chart = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.8), Inches(1.6),
                               Inches(8.4), Inches(5.2), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "y(t)"
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.text = "time (s)"

    # 8 click animation
    s = prs.slides.add_slide(layout_body)
    s.shapes.title.text = "Designing a controller"
    body = s.placeholders[1]
    steps = ["1. Model the system", "2. Choose the specifications",
             "3. Tune the controller", "4. Simulate and verify"]
    body.text_frame.text = steps[0]
    for t in steps[1:]:
        body.text_frame.add_paragraph().text = t
    appear_on_clicks(s, body.shape_id, [1, 2, 3])

    add_sections(prs, [("Introduction", [0, 1, 2]), ("Analysis", [3, 4, 5, 6]), ("Design", [7])])
    prs.save(path)
    return path


if __name__ == "__main__":
    print("wrote", build(os.path.join(HERE, "demo.pptx")))
